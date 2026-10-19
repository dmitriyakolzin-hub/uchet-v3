"""Выгрузка отчётов в Excel: заголовок, строка фильтров, таблица; суммы — числами с форматом денег."""

import io
from dataclasses import dataclass

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

MONEY_FORMAT = "#,##0.00;-#,##0.00;"  # ноль — пустая ячейка, как в отчёте на экране
QUANTITY_FORMAT = "#,##0.###;-#,##0.###;"


@dataclass
class XlsxColumn:
    title: str
    width: int = 16
    money: bool = False
    quantity: bool = False


def build(title: str, subtitle: str, columns: list[XlsxColumn], rows: list[tuple],
          total: list | None = None) -> bytes:
    """rows — (уровень вложенности, значения[, жирный]); уровень даёт отступ в первой колонке."""
    wb = Workbook()
    ws = wb.active
    ws.title = "".join("-" if ch in "[]:*?/\\" else ch for ch in title)[:31]  # запрещённые в имени листа
    ws.append([title])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([subtitle])
    ws.append([])
    ws.append([c.title for c in columns])
    header_row = ws.max_row
    for i, col in enumerate(columns, start=1):
        cell = ws.cell(row=header_row, column=i)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="E8EEF8")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = col.width
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)

    def put(values: list, level: int = 0, bold: bool = False) -> None:
        ws.append(values)  # Decimal записывается в файл как есть, без float
        r = ws.max_row
        for i, col in enumerate(columns, start=1):
            cell = ws.cell(row=r, column=i)
            if col.money:
                cell.number_format = MONEY_FORMAT
            elif col.quantity:
                cell.number_format = QUANTITY_FORMAT
            if bold:
                cell.font = Font(bold=True)
        if level:
            ws.cell(row=r, column=1).alignment = Alignment(indent=level)

    for level, values, *bold in rows:
        put(values, level, bool(bold and bold[0]))
    if total is not None:
        put(total, bold=True)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
