"""План бюджета (docs/TZ.md, раздел 13): ячейка сетки = одна плановая операция журнала.

Ячейка — версия + вид плана (pl — ОПиУ, cf — ДДС) + ЦФО + статья + месяц. Проводка ячейки строится
по шаблону статьи (Дт/Кт); ЦФО и статья ставятся в аналитику там, где её требует счёт. Плановые операции
отбираются той же select_lines(statuses=("plan",), budget_version_ids=...), что и факт.
"""

import datetime as dt
import io
from dataclasses import dataclass, field
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.money import round_money
from app.models import Account, Article, Budget, BudgetVersion, Cfo, LegalEntity, Operation, Posting, User
from app.models.journal import ANALYTIC_COLUMNS, analytic_column
from app.schemas.journal import AnalyticsIn, OperationIn, PostingIn
from app.services import journal
from app.services.excel_import import Column, RowError, RowProblem, _decimal, _read_rows, _text

MONTH_TITLES = ["Январь", "Февраль", "Март", "Апрель", "Май", "Июнь", "Июль", "Август", "Сентябрь", "Октябрь",
                "Ноябрь", "Декабрь"]
PL_ORDER = ["revenue", "cost", "selling", "admin", "other_income", "other_expense", "income_tax"]
PL_GROUPS = {"revenue": "Выручка", "cost": "Себестоимость продаж", "selling": "Коммерческие расходы",
             "admin": "Управленческие расходы", "other_income": "Прочие доходы", "other_expense": "Прочие расходы",
             "income_tax": "Налог на прибыль"}
CF_GROUPS = {("operating", "inflow"): "Операционная: поступления", ("operating", "outflow"): "Операционная: выплаты",
             ("investing", "inflow"): "Инвестиционная: поступления", ("investing", "outflow"): "Инвестиционная: выплаты",
             ("financing", "inflow"): "Финансовая: поступления", ("financing", "outflow"): "Финансовая: выплаты"}


class PlanError(ValueError):
    """Ячейку плана нельзя записать; текст — для пользователя."""


@dataclass
class Cell:
    article_id: int
    month: int
    amount: Decimal | None  # None или 0 — очистить ячейку


def month_key(year: int, month: int) -> str:
    return f"{year}-{month:02d}"


def plan_articles(db: Session, kind: str) -> list[tuple[Article, str]]:
    """Статьи сетки (с шаблоном проводки) и их группа, в порядке отчёта."""
    query = select(Article).where(Article.is_active, Article.debit_account_id.is_not(None))
    if kind == "pl":
        items = [(a, PL_GROUPS[a.pl_section]) for a in db.scalars(query.where(Article.is_pl))]
        order = {s: i for i, s in enumerate(PL_ORDER)}
        return sorted(items, key=lambda x: (order[x[0].pl_section], x[0].name.lower()))
    items = [(a, CF_GROUPS[(a.cash_flow_activity, a.cash_flow_direction)])
             for a in db.scalars(query.where(Article.is_cash_flow, ~Article.is_transfer))]
    keys = list(CF_GROUPS)
    return sorted(items, key=lambda x: (keys.index((x[0].cash_flow_activity, x[0].cash_flow_direction)),
                                        x[0].name.lower()))


def _legal_entity(db: Session, organization_id: int, cfo: Cfo | None) -> int:
    """Юрлицо плановой операции: юрлицо «Компании» в ветке ЦФО, иначе первое юрлицо организации."""
    if cfo is not None:
        ids = [int(i) for i in str(cfo.path).split(".")]
        for node in db.scalars(select(Cfo).where(Cfo.id.in_(ids))):
            if node.legal_entity_id:
                return node.legal_entity_id
    entity_id = db.scalar(select(LegalEntity.id).where(LegalEntity.organization_id == organization_id,
                                                       LegalEntity.is_active).order_by(LegalEntity.id).limit(1))
    if entity_id is None:
        raise PlanError("В организации нет юрлица — добавьте его в «Организации»")
    return entity_id


def _posting(db: Session, kind: str, article: Article, cfo_id: int | None, amount: Decimal) -> PostingIn:
    sides = {}
    for side, account_id in (("debit", article.debit_account_id), ("credit", article.credit_account_id)):
        required = {a.kind for a in db.get(Account, account_id).analytics}
        analytics = AnalyticsIn()
        if "cfo" in required and cfo_id is not None:
            analytics.cfo = cfo_id
        if "pl_item" in required and article.is_pl:
            analytics.pl_item = article.id
        if "cash_flow_item" in required and article.is_cash_flow:
            analytics.cash_flow_item = article.id
        sides[side] = analytics
    return PostingIn(debit_account_id=article.debit_account_id, credit_account_id=article.credit_account_id,
                     amount=amount, debit=sides["debit"], credit=sides["credit"])


def _cell_op(db: Session, version_id: int, kind: str, cfo_id: int | None, article_id: int, op_date: dt.date):
    query = select(Operation).where(Operation.budget_version_id == version_id, Operation.plan_kind == kind,
                                    Operation.article_id == article_id, Operation.op_date == op_date)
    query = query.where(Operation.plan_cfo_id.is_(None) if cfo_id is None else Operation.plan_cfo_id == cfo_id)
    return db.scalar(query)


def set_cells(db: Session, version: BudgetVersion, kind: str, cfo_id: int | None, cells: list[Cell], user: User) -> int:
    """Записать ячейки сетки; возвращает число изменённых ячеек. Версия — только черновик."""
    if version.status != "draft":
        raise PlanError("Менять можно только черновик версии; для пересмотра создайте новую версию")
    budget = db.get(Budget, version.budget_id)
    cfo = None
    if cfo_id is not None:
        cfo = db.get(Cfo, cfo_id)
        if cfo is None or cfo.organization_id != budget.organization_id:
            raise PlanError("ЦФО не найден в этой организации")
    if kind == "pl" and cfo is None:
        raise PlanError("План ОПиУ вводится по ЦФО — выберите ЦФО")
    allowed = {a.id: a for a, _ in plan_articles(db, kind)}
    entity_id = None
    changed = 0
    for cell in cells:
        article = allowed.get(cell.article_id)
        if article is None:
            raise PlanError("Статья не подходит для этого плана (нужен шаблон проводки и признак ДДС / ОПиУ)")
        if cell.month not in budget.months:
            raise PlanError(f"Месяц {cell.month} не входит в период бюджета «{budget.title}»")
        amount = round_money(cell.amount) if cell.amount is not None else None
        if amount is not None and amount < 0:
            raise PlanError(f"«{article.name}»: сумма плана не может быть отрицательной")
        op_date = dt.date(budget.year, cell.month, 1)
        op = _cell_op(db, version.id, kind, cfo_id, article.id, op_date)
        if not amount:
            if op is not None:
                db.delete(op)
                changed += 1
            continue
        if entity_id is None:
            entity_id = _legal_entity(db, budget.organization_id, cfo)
        if op is None:
            op = Operation(organization_id=budget.organization_id, number=0, status="plan", created_by=user.id,
                           budget_version_id=version.id, plan_kind=kind, plan_cfo_id=cfo_id)
        body = OperationIn(legal_entity_id=entity_id, op_date=op_date, article_id=article.id,
                           description=f"План: {article.name}", postings=[_posting(db, kind, article, cfo_id, amount)])
        journal.save_draft(db, op, body, user)
        db.add(op)
        changed += 1
    db.flush()
    return changed


def grid(db: Session, version: BudgetVersion, kind: str, cfo_id: int | None) -> dict:
    budget = db.get(Budget, version.budget_id)
    query = select(Operation).where(Operation.budget_version_id == version.id, Operation.plan_kind == kind)
    query = query.where(Operation.plan_cfo_id.is_(None) if cfo_id is None else Operation.plan_cfo_id == cfo_id)
    values: dict[int, dict[str, Decimal]] = {}
    for op in db.scalars(query):
        amount = sum((p.amount or Decimal(0) for p in op.postings), Decimal("0.00"))
        values.setdefault(op.article_id, {})[month_key(op.op_date.year, op.op_date.month)] = amount
    return {
        "months": [{"key": month_key(budget.year, m), "month": m, "title": MONTH_TITLES[m - 1]} for m in budget.months],
        "rows": [{"article_id": a.id, "name": a.name, "group": group, "values": values.get(a.id, {})}
                 for a, group in plan_articles(db, kind)],
    }


def copy_plan(db: Session, source: BudgetVersion, target: BudgetVersion, user: User) -> int:
    """Скопировать все плановые операции версии (ежемесячный пересмотр)."""
    count = 0
    for op in db.scalars(select(Operation).where(Operation.budget_version_id == source.id)):
        clone = Operation(
            organization_id=op.organization_id, legal_entity_id=op.legal_entity_id, number=0, op_date=op.op_date,
            article_id=op.article_id, description=op.description, status="plan", created_by=user.id,
            budget_version_id=target.id, plan_kind=op.plan_kind, plan_cfo_id=op.plan_cfo_id,
        )
        columns = ["line_no", "debit_account_id", "credit_account_id", "amount", "currency_id", "amount_rub",
                   "amount_rub_manual", "quantity"]
        columns += [analytic_column(side, kind) for side in ("debit", "credit") for kind in ANALYTIC_COLUMNS]
        clone.postings = [Posting(**{c: getattr(p, c) for c in columns}) for p in op.postings]
        db.add(clone)
        count += 1
    db.flush()
    return count


# --- Бюджет в Excel: скачать текущий план и загрузить исправленный ---


def cfo_labels(db: Session, organization_id: int, only: set[int] | None = None) -> dict[int, str]:
    """ЦФО → «Компания / Направление / …» (действующие; only — ограничение веткой)."""
    cfos = list(db.scalars(select(Cfo).where(Cfo.organization_id == organization_id)))
    names = {c.id: c.name for c in cfos}
    return {c.id: " / ".join(names[int(i)] for i in str(c.path).split("."))
            for c in cfos if c.is_active and (only is None or c.id in only)}


def plan_xlsx(db: Session, version: BudgetVersion, kind: str, cfo_ids: list[int | None]) -> bytes:
    budget = db.get(Budget, version.budget_id)
    labels = cfo_labels(db, budget.organization_id)
    wb = Workbook()
    ws = wb.active
    ws.title = "План ОПиУ" if kind == "pl" else "План ДДС"
    header = ["ЦФО", "Статья", *(MONTH_TITLES[m - 1] for m in budget.months)]
    ws.append(header)
    for i, title in enumerate(header, start=1):
        cell = ws.cell(row=1, column=i)
        cell.font, cell.fill = Font(bold=True), PatternFill("solid", fgColor="E8EEF8")
        ws.column_dimensions[get_column_letter(i)].width = 40 if i <= 2 else 14
    ws.freeze_panes = "C2"
    for cfo_id in cfo_ids:
        data = grid(db, version, kind, cfo_id)
        for row in data["rows"]:
            ws.append([labels.get(cfo_id, "") if cfo_id else "", row["name"],
                       *(row["values"].get(m["key"]) for m in data["months"])])
            for i in range(3, len(header) + 1):
                ws.cell(row=ws.max_row, column=i).number_format = "#,##0.00"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


@dataclass
class PlanImportResult:
    rows: int = 0
    changed: int = 0
    errors: list[RowError] = field(default_factory=list)


def import_plan(db: Session, version: BudgetVersion, kind: str, raw: bytes, user: User,
                branch: set[int] | None = None) -> PlanImportResult:
    """Строка файла: ЦФО + статья + суммы по месяцам; пустая ячейка или 0 — ячейка плана очищается.
    Строка с ошибкой не загружается (номер и причина — в результате)."""
    budget = db.get(Budget, version.budget_id)
    if version.status != "draft":
        raise PlanError("Загружать план можно только в черновик версии")
    columns = [Column("ЦФО", "cfo"), Column("Статья", "article", True)]
    columns += [Column(MONTH_TITLES[m - 1], f"m{m}") for m in budget.months]
    rows = _read_rows(raw, columns)
    labels = cfo_labels(db, budget.organization_id)
    by_label = {v.lower(): k for k, v in labels.items()}
    by_name: dict[str, list[int]] = {}
    for cfo_id, label in labels.items():
        by_name.setdefault(label.split(" / ")[-1].lower(), []).append(cfo_id)
    articles = {a.name.lower(): a for a, _ in plan_articles(db, kind)}
    result = PlanImportResult(rows=len(rows))
    for number, row in rows:
        try:
            cfo_text = _text(row.get("cfo"))
            cfo_id = None
            if cfo_text:
                cfo_id = by_label.get(cfo_text.lower())
                if cfo_id is None:
                    found = by_name.get(cfo_text.lower(), [])
                    if len(found) != 1:
                        raise RowProblem(f"ЦФО «{cfo_text}» не найден" if not found
                                         else f"ЦФО «{cfo_text}»: несколько с таким названием — укажите путь")
                    cfo_id = found[0]
            if kind == "pl" and cfo_id is None:
                raise RowProblem("Не указан ЦФО")
            if branch is not None and cfo_id not in branch:
                raise RowProblem("ЦФО вне вашей ветки")
            article = articles.get(_text(row.get("article")).lower())
            if article is None:
                raise RowProblem(f"Статьи «{_text(row.get('article'))}» нет в плане {'ОПиУ' if kind == 'pl' else 'ДДС'}")
            cells = []
            for m in budget.months:
                value = row.get(f"m{m}")
                amount = None if _text(value) in ("", "0") else _decimal(value, MONTH_TITLES[m - 1], 2)
                cells.append(Cell(article.id, m, amount))
            with db.begin_nested():
                result.changed += set_cells(db, version, kind, cfo_id, cells, user)
        except (RowProblem, PlanError) as exc:
            result.errors.append(RowError(number, str(exc)))
    return result
