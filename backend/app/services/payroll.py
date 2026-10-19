"""Зарплата (docs/TZ.md, раздел 14)."""

import calendar
import datetime as dt
import io
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.core import payroll_rates as rates
from app.core.money import round_money
from app.models import (
    Account,
    Article,
    BankAccount,
    Currency,
    Employee,
    EmployeeSalary,
    LegalEntity,
    Operation,
    PayoutLine,
    PayoutSheet,
    PayrollLine,
    PayrollRun,
    User,
)
from app.schemas.journal import AnalyticsIn, OperationIn, PostingIn
from app.services import journal
from app.services.excel_import import Column, RowError, RowProblem, _decimal, _read_rows, _text
from app.services.journal import JournalError


def salary_on(db: Session, employee_id: int, on_date: dt.date) -> Decimal | None:
    """Оклад на дату: последний с датой начала не позже этой даты."""
    return db.scalar(
        select(EmployeeSalary.amount)
        .where(EmployeeSalary.employee_id == employee_id, EmployeeSalary.valid_from <= on_date)
        .order_by(EmployeeSalary.valid_from.desc())
        .limit(1)
    )


# --- Начисление зарплаты за месяц ---

INSURANCE_ARTICLE = "Страховые взносы с оплаты труда"
MONTHS_GENITIVE = ["январь", "февраль", "март", "апрель", "май", "июнь", "июль", "август", "сентябрь",
                   "октябрь", "ноябрь", "декабрь"]
ZERO = Decimal("0.00")


class PayrollError(ValueError):
    """Начисление нельзя выполнить; текст — для пользователя."""


def month_bounds(year: int, month: int) -> tuple[dt.date, dt.date]:
    return dt.date(year, month, 1), dt.date(year, month, calendar.monthrange(year, month)[1])


def default_norm(year: int, month: int) -> int:
    return rates.working_days(*month_bounds(year, month))


def _employee_days(employee: Employee, run: PayrollRun) -> Decimal:
    start, end = month_bounds(run.year, run.month)
    first = max(start, employee.hired_on) if employee.hired_on else start
    last = min(end, employee.fired_on) if employee.fired_on else end
    if (first, last) == (start, end):
        return Decimal(run.norm_days)
    return Decimal(rates.working_days(first, last)) if first <= last else Decimal(0)


def fill(db: Session, run: PayrollRun) -> int:
    """Добавить сотрудников юрлица, работавших в месяце; оклад — на последний день месяца."""
    start, end = month_bounds(run.year, run.month)
    present = set(db.scalars(select(PayrollLine.employee_id).where(PayrollLine.run_id == run.id)))
    employees = db.scalars(select(Employee).where(
        Employee.organization_id == run.organization_id, Employee.legal_entity_id == run.legal_entity_id,
        Employee.is_active, (Employee.hired_on.is_(None)) | (Employee.hired_on <= end),
        (Employee.fired_on.is_(None)) | (Employee.fired_on >= start),
    ).order_by(Employee.full_name))
    added = 0
    for employee in employees:
        if employee.id in present:
            continue
        db.add(PayrollLine(run_id=run.id, employee_id=employee.id, salary=salary_on(db, employee.id, end) or ZERO,
                           work_days=_employee_days(employee, run)))
        added += 1
    db.flush()
    recalc(db, run)
    return added


def _ytd_before(db: Session, run: PayrollRun, employee_id: int) -> tuple[Decimal, Decimal]:
    """Доход и НДФЛ сотрудника по проведённым начислениям этого юрлица с начала года до месяца run."""
    income, tax = db.execute(
        select(func.coalesce(func.sum(PayrollLine.accrued), 0), func.coalesce(func.sum(PayrollLine.ndfl), 0))
        .join(PayrollRun, PayrollRun.id == PayrollLine.run_id)
        .where(PayrollRun.legal_entity_id == run.legal_entity_id, PayrollRun.year == run.year,
               PayrollRun.month < run.month, PayrollRun.status == "posted", PayrollLine.employee_id == employee_id)
    ).one()
    return Decimal(income), Decimal(tax)


def recalc(db: Session, run: PayrollRun) -> None:
    tariff = db.get(LegalEntity, run.legal_entity_id).insurance_tariff
    for line in db.scalars(select(PayrollLine).where(PayrollLine.run_id == run.id)):
        line.salary_amount = round_money(line.salary * line.work_days / run.norm_days)
        line.accrued = line.salary_amount + line.bonus + line.vacation + line.sick + line.other_accrual
        income_before, tax_before = _ytd_before(db, run, line.employee_id)
        line.ndfl = max(ZERO, rates.ndfl_total(income_before + line.accrued) - tax_before) + ZERO
        # пособие по болезни взносами не облагается
        line.insurance = rates.insurance(tariff, run.year, line.accrued - line.sick, income_before)
        line.net = line.accrued - line.ndfl - line.other_deduction
    db.flush()


def _posting(debit, credit, amount, debit_an=None, credit_an=None) -> PostingIn:
    return PostingIn(debit_account_id=debit, credit_account_id=credit, amount=amount,
                     debit=debit_an or AnalyticsIn(), credit=credit_an or AnalyticsIn())


def post(db: Session, run: PayrollRun, user: User) -> Operation:
    """Провести: одна операция журнала на последний день месяца (проводки по каждому сотруднику)."""
    if run.status != "draft":
        raise PayrollError("Начисление уже проведено")
    lines = list(db.scalars(select(PayrollLine).where(PayrollLine.run_id == run.id)))
    if not lines:
        raise PayrollError("В начислении нет сотрудников — нажмите «Заполнить»")
    recalc(db, run)
    accounts = {a.code: a.id for a in db.scalars(select(Account))}
    insurance_article = db.scalar(select(Article).where(Article.name == INSURANCE_ARTICLE))
    if insurance_article is None:
        raise PayrollError(f"Нет статьи «{INSURANCE_ARTICLE}» — восстановите её в справочнике статей")
    problems, postings = [], []
    for line in lines:
        employee = db.get(Employee, line.employee_id)
        article = db.get(Article, employee.pl_item_id) if employee.pl_item_id else None
        if employee.cfo_id is None or article is None or article.debit_account_id is None:
            problems.append(employee.full_name)
            continue
        cost = article.debit_account_id
        cost_an = AnalyticsIn(cfo=employee.cfo_id, pl_item=article.id)
        staff = AnalyticsIn(employee=employee.id)
        if line.accrued > 0:
            postings.append(_posting(cost, accounts["70"], line.accrued, cost_an, staff))
        if line.ndfl > 0:
            postings.append(_posting(accounts["70"], accounts["68.01"], line.ndfl, staff))
        if line.insurance > 0:
            postings.append(_posting(cost, accounts["69"], line.insurance,
                                     AnalyticsIn(cfo=employee.cfo_id, pl_item=insurance_article.id)))
        if line.other_deduction > 0:
            postings.append(_posting(accounts["70"], accounts["73"], line.other_deduction, staff, staff))
    if problems:
        raise PayrollError("У сотрудников не заполнены ЦФО или статья затрат (с шаблоном проводки): "
                           + ", ".join(problems))
    if not postings:
        raise PayrollError("Начислять нечего — все суммы нулевые")
    _, end = month_bounds(run.year, run.month)
    body = OperationIn(legal_entity_id=run.legal_entity_id, op_date=end,
                       description=f"Начисление зарплаты за {MONTHS_GENITIVE[run.month - 1]} {run.year}",
                       postings=postings)
    op = Operation(organization_id=run.organization_id, created_by=user.id, status="draft")
    try:
        journal.save_draft(db, op, body, user)
        op.number = journal.next_number(db, run.organization_id)
        db.add(op)
        db.flush()
        journal._event(db, op, user, "created")
        journal.post(db, op, user)
    except JournalError as exc:
        raise PayrollError(f"Проводки не прошли проверку журнала: {exc}") from None
    run.status, run.operation_id = "posted", op.id
    db.flush()
    return op


def unpost(db: Session, run: PayrollRun) -> None:
    """Отменить проведение: операция журнала удаляется, начисление снова черновик."""
    if run.status != "posted":
        raise PayrollError("Начисление не проведено")
    later = db.scalar(select(PayrollRun.id).where(
        PayrollRun.legal_entity_id == run.legal_entity_id, PayrollRun.year == run.year,
        PayrollRun.month > run.month, PayrollRun.status == "posted").limit(1))
    if later:
        raise PayrollError("Сначала отмените проведение начислений следующих месяцев — от них зависит НДФЛ")
    if run.operation_id:
        op = db.get(Operation, run.operation_id)
        if op is not None:
            db.delete(op)
    run.status, run.operation_id = "draft", None
    db.flush()


# --- Начисление в Excel: скачать строки, поправить дни и суммы, загрузить ---

LINE_COLUMNS = [
    ("personnel_number", "Таб. №"), ("employee", "Сотрудник"), ("work_days", "Отработано дней"),
    ("bonus", "Премия"), ("vacation", "Отпускные"), ("sick", "Больничные"),
    ("other_accrual", "Прочие начисления"), ("other_deduction", "Прочие удержания"),
]


def lines_xlsx(db: Session, run: PayrollRun) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = f"Зарплата {run.month:02d}.{run.year}"
    ws.append([title for _, title in LINE_COLUMNS])
    for i in range(1, len(LINE_COLUMNS) + 1):
        ws.cell(row=1, column=i).font = Font(bold=True)
        ws.cell(row=1, column=i).fill = PatternFill("solid", fgColor="E8EEF8")
        ws.column_dimensions[chr(64 + i)].width = 30 if i == 2 else 16
    rows = db.execute(select(PayrollLine, Employee).join(Employee, Employee.id == PayrollLine.employee_id)
                      .where(PayrollLine.run_id == run.id).order_by(Employee.full_name)).all()
    for line, employee in rows:
        ws.append([employee.personnel_number or None, employee.full_name, line.work_days, line.bonus, line.vacation,
                   line.sick, line.other_accrual, line.other_deduction])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def import_lines(db: Session, run: PayrollRun, raw: bytes) -> tuple[int, int, list]:
    """Строка файла: сотрудник (таб. № или ФИО) + дни и суммы; пустая сумма — 0. Сотрудник юрлица,
    которого нет в начислении, добавляется. Возвращает (строк, изменено, ошибки)."""
    columns = [Column(title, key, key == "employee") for key, title in LINE_COLUMNS]
    rows = _read_rows(raw, columns)
    employees = list(db.scalars(select(Employee).where(Employee.organization_id == run.organization_id,
                                                       Employee.legal_entity_id == run.legal_entity_id)))
    lines = {line.employee_id: line for line in db.scalars(select(PayrollLine).where(PayrollLine.run_id == run.id))}
    changed, errors = 0, []
    for number, row in rows:
        try:
            tab, name = _text(row.get("personnel_number")), _text(row.get("employee"))
            found = [e for e in employees if tab and e.personnel_number == tab] or \
                [e for e in employees if e.full_name.strip().lower() == name.lower()]
            if len(found) != 1:
                raise RowProblem(f"Сотрудник «{name or tab}» юрлица не найден" if not found
                                 else f"«{name}»: несколько сотрудников — укажите табельный №")
            employee = found[0]
            values = {}
            for key, title in LINE_COLUMNS[2:]:
                value = row.get(key)
                if key == "work_days":
                    if _text(value) == "":
                        continue
                    days = _decimal(value, title, 2) if _text(value) != "0" else ZERO
                    if days > 31:
                        raise RowProblem("«Отработано дней»: не больше 31")
                    values[key] = days
                else:
                    values[key] = ZERO if _text(value) in ("", "0") else _decimal(value, title, 2)
            line = lines.get(employee.id)
            if line is None:
                line = PayrollLine(run_id=run.id, employee_id=employee.id,
                                   salary=salary_on(db, employee.id, month_bounds(run.year, run.month)[1]) or ZERO,
                                   work_days=_employee_days(employee, run))
                db.add(line)
                lines[employee.id] = line
            for key, value in values.items():
                setattr(line, key, value)
            changed += 1
        except RowProblem as exc:
            errors.append(RowError(number, str(exc)))
    db.flush()
    recalc(db, run)
    return len(rows), changed, errors


# --- Ведомости на выплату: аванс и зарплата, оплата через банк, подтверждение выпиской ---

SALARY_PAYMENT_ARTICLE = "Выплата заработной платы"
PAYMENT_WORDS = ("заработн", "зарплат", "аванс", "реестр", "оплата труда")


def sheet_total(db: Session, sheet: PayoutSheet) -> Decimal:
    return db.scalar(select(func.coalesce(func.sum(PayoutLine.amount), 0)).where(PayoutLine.sheet_id == sheet.id)) + ZERO


def _advances_paid(db: Session, sheet: PayoutSheet, employee_id: int) -> Decimal:
    """Авансы сотруднику за месяц по ведомостям, отправленным в банк или оплаченным."""
    return db.scalar(
        select(func.coalesce(func.sum(PayoutLine.amount), 0)).join(PayoutSheet, PayoutSheet.id == PayoutLine.sheet_id)
        .where(PayoutSheet.legal_entity_id == sheet.legal_entity_id, PayoutSheet.year == sheet.year,
               PayoutSheet.month == sheet.month, PayoutSheet.kind == "advance",
               PayoutSheet.status.in_(("sent", "paid")), PayoutLine.employee_id == employee_id)
    ) + ZERO


def fill_sheet(db: Session, sheet: PayoutSheet, advance_percent: Decimal = Decimal(50)) -> int:
    """Аванс — доля оклада сотрудников, работающих в месяце; зарплата — «к выплате» из проведённого
    начисления минус авансы месяца."""
    db.query(PayoutLine).filter(PayoutLine.sheet_id == sheet.id).delete()
    amounts: dict[int, Decimal] = {}
    if sheet.kind == "advance":
        start, end = month_bounds(sheet.year, sheet.month)
        employees = db.scalars(select(Employee).where(
            Employee.organization_id == sheet.organization_id, Employee.legal_entity_id == sheet.legal_entity_id,
            Employee.is_active, (Employee.hired_on.is_(None)) | (Employee.hired_on <= end),
            (Employee.fired_on.is_(None)) | (Employee.fired_on >= start)))
        for employee in employees:
            salary = salary_on(db, employee.id, sheet.pay_date) or ZERO
            amounts[employee.id] = round_money(salary * advance_percent / 100)
    else:
        run = db.scalar(select(PayrollRun).where(
            PayrollRun.legal_entity_id == sheet.legal_entity_id, PayrollRun.year == sheet.year,
            PayrollRun.month == sheet.month, PayrollRun.status == "posted"))
        if run is None:
            raise PayrollError("Сначала проведите начисление зарплаты за этот месяц")
        for line in db.scalars(select(PayrollLine).where(PayrollLine.run_id == run.id)):
            amounts[line.employee_id] = line.net - _advances_paid(db, sheet, line.employee_id)
    count = 0
    for employee_id, amount in amounts.items():
        if amount > 0:
            db.add(PayoutLine(sheet_id=sheet.id, employee_id=employee_id, amount=amount))
            count += 1
    db.flush()
    return count


def pay_sheet(db: Session, sheet: PayoutSheet, user: User, paid_on: dt.date, *, import_key: str | None = None,
              confirmed: bool = False) -> Operation:
    """Оплата: проведённая операция Дт 70 (сотрудник) — Кт 51 (расчётный счёт, статья выплаты)."""
    if sheet.status != "sent":
        raise PayrollError("Оплатить можно ведомость, отправленную в банк")
    article = db.scalar(select(Article).where(Article.name == SALARY_PAYMENT_ARTICLE))
    if article is None:
        raise PayrollError(f"Нет статьи «{SALARY_PAYMENT_ARTICLE}» — восстановите её в справочнике статей")
    bank = db.get(BankAccount, sheet.bank_account_id)
    money_code = "51" if db.get(Currency, bank.currency_id).code == "RUB" else "52"
    accounts = {a.code: a.id for a in db.scalars(select(Account).where(Account.code.in_(("70", money_code))))}
    postings = [
        _posting(accounts["70"], accounts[money_code], line.amount, AnalyticsIn(employee=line.employee_id),
                 AnalyticsIn(bank_account=bank.id, cash_flow_item=article.id))
        for line in db.scalars(select(PayoutLine).where(PayoutLine.sheet_id == sheet.id))
    ]
    kind = "аванса" if sheet.kind == "advance" else "зарплаты"
    body = OperationIn(legal_entity_id=sheet.legal_entity_id, op_date=paid_on, article_id=article.id,
                       description=f"Выплата {kind} за {MONTHS_GENITIVE[sheet.month - 1]} {sheet.year} "
                                   f"по ведомости № {sheet.number}", postings=postings)
    op = Operation(organization_id=sheet.organization_id, created_by=user.id, status="draft", import_key=import_key)
    try:
        journal.save_draft(db, op, body, user)
        op.number = journal.next_number(db, sheet.organization_id)
        db.add(op)
        db.flush()
        journal._event(db, op, user, "created")
        journal.post(db, op, user)
    except JournalError as exc:
        raise PayrollError(f"Оплата не прошла проверку журнала: {exc}") from None
    sheet.status, sheet.operation_id, sheet.paid_on, sheet.confirmed_by_statement = "paid", op.id, paid_on, confirmed
    db.flush()
    return op


def unpay_sheet(db: Session, sheet: PayoutSheet) -> None:
    if sheet.status != "paid":
        raise PayrollError("Ведомость не оплачена")
    if sheet.operation_id and (op := db.get(Operation, sheet.operation_id)) is not None:
        db.delete(op)
    sheet.status, sheet.operation_id, sheet.paid_on, sheet.confirmed_by_statement = "sent", None, None, False
    db.flush()


def match_sheet(db: Session, organization_id: int, bank_account_id: int, amount: Decimal,
                purpose: str) -> PayoutSheet | None:
    """Отправленная в банк ведомость, которую оплатил этот платёж: тот же счёт, та же сумма, назначение о зарплате."""
    if not any(word in purpose.lower() for word in PAYMENT_WORDS):
        return None
    for sheet in db.scalars(select(PayoutSheet).where(
            PayoutSheet.organization_id == organization_id, PayoutSheet.bank_account_id == bank_account_id,
            PayoutSheet.status == "sent").order_by(PayoutSheet.pay_date, PayoutSheet.number)):
        if sheet_total(db, sheet) == amount:
            return sheet
    return None


def registry_xlsx(db: Session, sheet: PayoutSheet) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = f"Ведомость {sheet.number}"
    kind = "Аванс" if sheet.kind == "advance" else "Зарплата"
    ws.append([f"{kind} за {MONTHS_GENITIVE[sheet.month - 1]} {sheet.year}, ведомость № {sheet.number}, "
               f"дата выплаты {sheet.pay_date:%d.%m.%Y}"])
    ws.append([])
    ws.append(["Таб. №", "ФИО", "ИНН", "Сумма"])
    for i in range(1, 5):
        ws.cell(row=3, column=i).font = Font(bold=True)
        ws.cell(row=3, column=i).fill = PatternFill("solid", fgColor="E8EEF8")
    ws.column_dimensions["B"].width = 36
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["D"].width = 16
    rows = db.execute(select(PayoutLine, Employee).join(Employee, Employee.id == PayoutLine.employee_id)
                      .where(PayoutLine.sheet_id == sheet.id).order_by(Employee.full_name)).all()
    for line, employee in rows:
        ws.append([employee.personnel_number or None, employee.full_name, employee.inn, line.amount])
        ws.cell(row=ws.max_row, column=4).number_format = "#,##0.00"
    ws.append([None, "Итого", None, sheet_total(db, sheet)])
    ws.cell(row=ws.max_row, column=2).font = Font(bold=True)
    ws.cell(row=ws.max_row, column=4).font = Font(bold=True)
    ws.cell(row=ws.max_row, column=4).number_format = "#,##0.00"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# --- Средний заработок: отпускные и больничные (docs/TZ.md, раздел 15, «Детали 5.4») ---

DAYS_PER_MONTH = Decimal("29.3")
SICK_PERCENT = {"lt6m": Decimal(100), "lt5": Decimal(60), "5to8": Decimal(80), "ge8": Decimal(100)}
EMPLOYER_SICK_DAYS = 3


def _ru(value: Decimal) -> str:
    return f"{value:,.2f}".replace(",", " ").replace(".", ",")


def _posted_lines(db: Session, employee_id: int, first: tuple[int, int], last: tuple[int, int]):
    """Строки проведённых начислений сотрудника за месяцы first..last включительно, с нормой дней."""
    lo, hi = first[0] * 12 + first[1], last[0] * 12 + last[1]
    return db.execute(select(PayrollLine, PayrollRun).join(PayrollRun, PayrollRun.id == PayrollLine.run_id).where(
        PayrollLine.employee_id == employee_id, PayrollRun.status == "posted",
        PayrollRun.year * 12 + PayrollRun.month >= lo, PayrollRun.year * 12 + PayrollRun.month <= hi)).all()


def calc_vacation(db: Session, line: PayrollLine, days: int) -> None:
    """Отпускные: выплаты за 12 месяцев до месяца начисления / (29,3 × месяцы) × календарные дни отпуска."""
    run = db.get(PayrollRun, line.run_id)
    end = (run.year, run.month - 1) if run.month > 1 else (run.year - 1, 12)
    start = (end[0] - 1, end[1] + 1) if end[1] < 12 else (end[0], 1)
    earnings, base_days, months = ZERO, Decimal(0), 0
    for past, past_run in _posted_lines(db, line.employee_id, start, end):
        earnings += past.salary_amount + past.bonus + past.other_accrual
        share = min(past.work_days / past_run.norm_days, Decimal(1)) if past_run.norm_days else Decimal(0)
        base_days += DAYS_PER_MONTH * share
        months += 1
    if base_days > 0:
        average = round_money(earnings / base_days)
        how = (f"выплаты за 12 мес. ({months} мес. в программе) {_ru(earnings)} ₽ / "
               f"{str(base_days.quantize(Decimal('0.01'))).replace('.', ',')} дн.")
    else:
        average = round_money(line.salary / DAYS_PER_MONTH)
        how = f"истории начислений нет — оклад {_ru(line.salary)} ₽ / 29,3"
    line.vacation = round_money(average * days)
    line.vacation_note = f"Отпуск {days} дн. × средний {_ru(average)} ₽ ({how}) = {_ru(line.vacation)} ₽"
    recalc(db, run)


def calc_sick(db: Session, line: PayrollLine, days: int, experience: str,
              earnings_by_year: dict[int, Decimal] | None = None) -> None:
    """Больничный: заработок за 2 предыдущих года (не больше предельной базы) / 730 × процент по стажу;
    работодатель платит первые 3 дня — они и попадают в начисление."""
    run = db.get(PayrollRun, line.run_id)
    years = [run.year - 2, run.year - 1]
    parts, total = [], ZERO
    for year in years:
        if earnings_by_year and earnings_by_year.get(year) is not None:
            amount, source = earnings_by_year[year], "введено"
        else:
            amount = sum((p.salary_amount + p.bonus + p.vacation + p.other_accrual
                          for p, _ in _posted_lines(db, line.employee_id, (year, 1), (year, 12))), ZERO)
            source = "по программе"
        capped = min(amount, rates.base_limit(year))
        total += capped
        parts.append(f"{year}: {_ru(capped)} ₽ ({source})")
    minimum = round_money(rates.mrot(run.year) * 24 / 730)
    maximum = round_money((rates.base_limit(years[0]) + rates.base_limit(years[1])) / 730)
    average = minimum if experience == "lt6m" else min(max(round_money(total / 730), minimum), maximum)
    daily = round_money(average * SICK_PERCENT[experience] / 100)
    employer_days = min(days, EMPLOYER_SICK_DAYS)
    line.sick = round_money(daily * employer_days)
    fund = round_money(daily * (days - employer_days))
    line.sick_note = (f"Больничный {days} дн.: средний {_ru(average)} ₽ ({'; '.join(parts)}; / 730, "
                      f"мин. {_ru(minimum)}) × {SICK_PERCENT[experience]}% = {_ru(daily)} ₽ в день; "
                      f"работодатель {employer_days} дн. = {_ru(line.sick)} ₽, СФР {days - employer_days} дн. = {_ru(fund)} ₽")
    recalc(db, run)
