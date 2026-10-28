"""Загрузка из Excel (docs/TZ.md, раздел 11, часть 1.7б): операции по шаблону и список контрагентов.

Шаблон операций: строка = проводка; строки с одинаковым значением в колонке «Операция» собираются
в одну операцию, без значения — каждая строка отдельная операция. Операции становятся черновиками
через journal.save_draft (те же проверки, что при ручном вводе). Строка с ошибкой (не найден счёт,
контрагент и т.п.) не загружается — её номер и причина возвращаются в результате; остальные строки
загружаются. Повторная загрузка того же файла не создаёт дублей (import_key).
"""

import datetime as dt
import hashlib
import io
import re
import zipfile
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.analytics import ANALYTIC_TITLES, OPTIONAL_ANALYTICS, AnalyticKind
from app.core.money import round_money
from app.models import (
    Account,
    Article,
    BankAccount,
    Cfo,
    Contract,
    Counterparty,
    Currency,
    Employee,
    Executor,
    FixedAsset,
    LegalEntity,
    Nomenclature,
    Operation,
    User,
    Warehouse,
)
from app.schemas.journal import AnalyticsIn, OperationIn, PostingIn
from app.schemas.reference import check_inn_for_kind
from app.services import journal
from app.services.journal import OPENING_ACCOUNT, JournalError

MAX_ROWS = 5000
XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


class ExcelError(ValueError):
    """Файл нельзя разобрать целиком; текст — для пользователя."""


@dataclass
class Column:
    title: str
    key: str
    required: bool = False
    width: int = 16
    text: bool = False  # ячейка в шаблоне — текстовая (номера счетов, ИНН не превращаются в числа)
    hint: str = ""


OPERATION_COLUMNS = [
    Column("Дата", "date", True, 12, hint="ДД.ММ.ГГГГ"),
    Column("Операция", "group", False, 11, True,
           "Необязательно. Строки с одинаковым значением (например, 1) станут одной операцией из нескольких проводок"),
    Column("Юрлицо (ИНН)", "entity", False, 15, True, "Можно не заполнять, если юрлицо в организации одно"),
    Column("Статья", "article", False, 32, hint="Статья ДДС/ОПиУ из справочника; ставится в аналитику, где она нужна"),
    Column("Счёт Дт", "debit", True, 10, True, "Код счёта, например 60.01"),
    Column("Счёт Кт", "credit", True, 10, True, "Код счёта, например 51"),
    Column("Сумма", "amount", True, 14, hint="Число, не больше 2 знаков после запятой"),
    Column("Валюта", "currency", False, 8, True, "Код валюты; пусто — RUB"),
    Column("Количество", "quantity", False, 11, hint="Только для количественных счетов (10, 41, 43)"),
    Column("Контрагент (ИНН или наименование)", "counterparty", False, 28, True),
    Column("Договор №", "contract", False, 14, True, "Номер договора с этим контрагентом"),
    Column("Расчётный счёт (номер)", "bank_account", False, 24, True, "20 цифр"),
    Column("Сотрудник (таб. № или ФИО)", "employee", False, 24, True),
    Column("Номенклатура (артикул или наименование)", "nomenclature", False, 28, True),
    Column("Склад", "warehouse", False, 18, True),
    Column("Основное средство (инв. № или наименование)", "fixed_asset", False, 28, True,
           "Для счетов 01, 02, 08"),
    Column("ЦФО", "cfo", False, 22, True, "Код ЦФО, путь «Компания / … / Магазин» или название"),
    Column("ЦФУ", "executor", False, 22, True, "Код ЦФУ, «Отдел / Исполнитель» или название; ставится там же, где ЦФО"),
    Column("Описание", "description", False, 40),
]

COUNTERPARTY_COLUMNS = [
    Column("Наименование", "name", True, 32),
    Column("Полное наименование", "full_name", False, 40),
    Column("Вид", "kind", False, 14, hint="Организация / ИП / Физлицо; пусто — по длине ИНН"),
    Column("ИНН", "inn", False, 14, True, "10 цифр — организация, 12 — ИП или физлицо"),
    Column("КПП", "kpp", False, 11, True),
    Column("Покупатель", "is_customer", False, 11, hint="да / нет"),
    Column("Поставщик", "is_supplier", False, 11, hint="да / нет"),
]

# Колонки аналитики шаблона → вид аналитики (статья ставится отдельно — по признакам статьи)
ANALYTIC_KEYS = ["counterparty", "contract", "bank_account", "employee", "nomenclature", "warehouse", "fixed_asset",
                 "cfo", "executor"]
KIND_TITLES = {"организация": "legal", "юрлицо": "legal", "юл": "legal",
               "ип": "entrepreneur", "индивидуальный предприниматель": "entrepreneur",
               "физлицо": "person", "физ. лицо": "person", "физическое лицо": "person", "фл": "person"}
YES = {"да", "yes", "1", "+", "true", "истина", "x", "х"}
NO = {"", "нет", "no", "0", "-", "false", "ложь"}


@dataclass
class RowError:
    row: int
    error: str


@dataclass
class OperationsResult:
    rows: int = 0  # строк с данными
    created: list[int] = field(default_factory=list)
    incomplete: list[int] = field(default_factory=list)
    skipped_duplicates: int = 0
    errors: list[RowError] = field(default_factory=list)


@dataclass
class CounterpartiesResult:
    rows: int = 0
    created: list[str] = field(default_factory=list)
    skipped_existing: int = 0
    errors: list[RowError] = field(default_factory=list)


class RowProblem(ValueError):
    pass


# --- Шаблоны ---


def _sheet_with_header(ws, columns: list[Column]) -> None:
    bold, fill = Font(bold=True), PatternFill("solid", fgColor="E8EEF8")
    for i, col in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=i, value=col.title + ("*" if col.required else ""))
        cell.font, cell.fill = bold, fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        letter = get_column_letter(i)
        ws.column_dimensions[letter].width = col.width
        if col.text:
            for row in range(2, 1001):
                ws.cell(row=row, column=i).number_format = "@"
    ws.freeze_panes = "A2"


def _help_sheet(wb: Workbook, columns: list[Column], intro: list[str]) -> None:
    ws = wb.create_sheet("Как заполнять")
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 100
    for line in intro:
        ws.append([line])
    ws.append([])
    ws.append(["Колонка", "Что указать"])
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
    ws.cell(row=ws.max_row, column=2).font = Font(bold=True)
    for col in columns:
        ws.append([col.title + (" (обязательно)" if col.required else ""), col.hint])


def _to_bytes(wb: Workbook) -> bytes:
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def leaf_articles(db: Session, organization_id: int, **conditions) -> list[Article]:
    """Действующие статьи организации без подстатей — только их ставят в операции."""
    groups = select(Article.parent_id).where(Article.parent_id.is_not(None))
    query = select(Article).where(Article.organization_id == organization_id, Article.is_active, Article.id.not_in(groups))
    for name, value in conditions.items():
        query = query.where(getattr(Article, name) == value)
    return list(db.scalars(query))


def article_label(article: Article) -> str:
    """Как статья пишется в файле: «код название» или название."""
    return f"{article.code} {article.name}" if article.code else article.name


class ArticleLookup:
    """Статья по ячейке: код, «код название» или название (если такое название одно)."""

    def __init__(self, items: list[Article]):
        self.by_code = {a.code.rstrip("."): a for a in items if a.code}
        self.by_name: dict[str, list[Article]] = {}
        for a in items:
            self.by_name.setdefault(a.name.strip().lower(), []).append(a)

    def find(self, text: str, title: str) -> Article:
        head, _, rest = text.partition(" ")
        article = self.by_code.get(text.rstrip(".")) or (self.by_code.get(head.rstrip(".")) if rest else None)
        if article is not None:
            return article
        found = self.by_name.get(text.strip().lower(), [])
        if len(found) > 1:
            raise RowProblem(f"«{title}»: статей «{text}» несколько — укажите код статьи")
        if not found:
            raise RowProblem(f"«{title}»: статьи «{text}» нет в справочнике (или это группа — укажите подстатью)")
        return found[0]


class TreeLookup:
    """ЦФО или ЦФУ по ячейке: код, путь («Компания / … / Магазин», «Отдел / Исполнитель») или название, если оно одно."""

    def __init__(self, items: list, labels: dict[int, str], title: str):
        self.title = title
        self.by_code = {x.code.strip().lower(): x.id for x in items if x.code}
        self.by_text: dict[str, list[int]] = {}  # путь и название
        for x in items:
            for text in {labels[x.id].lower(), x.name.strip().lower()}:
                self.by_text.setdefault(text, []).append(x.id)

    def find(self, text: str) -> int:
        low = text.strip().lower()
        if low in self.by_code:
            return self.by_code[low]
        ids = self.by_text.get(low, [])
        if not ids:
            raise RowProblem(f"{self.title} «{text}» не найден в справочнике организации")
        if len(ids) > 1:
            raise RowProblem(f"{self.title} «{text}»: несколько с таким названием — укажите код или путь")
        return ids[0]


def cfo_lookup(db: Session, organization_id: int, active_only: bool = True) -> TreeLookup:
    cfos = list(db.scalars(select(Cfo).where(Cfo.organization_id == organization_id)))
    names = {c.id: c.name for c in cfos}
    labels = {c.id: " / ".join(names[int(i)] for i in str(c.path).split(".")) for c in cfos}
    return TreeLookup([c for c in cfos if c.is_active or not active_only], labels, "ЦФО")


def executor_lookup(db: Session, organization_id: int, active_only: bool = True) -> TreeLookup:
    items = list(db.scalars(select(Executor).where(Executor.organization_id == organization_id)))
    names = {e.id: e.name for e in items}
    labels = {e.id: f"{names[e.parent_id]} / {e.name}" if e.parent_id else e.name for e in items}
    return TreeLookup([e for e in items if e.is_active or not active_only], labels, "ЦФУ")


def operations_template(db: Session, organization_id: int) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Операции"
    _sheet_with_header(ws, OPERATION_COLUMNS)

    # Списки для подсказок и выпадающих списков: счета без субсчетов и действующие статьи
    parents = set(db.scalars(select(Account.parent_id).where(Account.parent_id.is_not(None))))
    accounts = [a for a in db.scalars(select(Account).where(Account.is_active).order_by(Account.code))
                if a.id not in parents and a.code != OPENING_ACCOUNT]
    articles = sorted(leaf_articles(db, organization_id), key=lambda a: (a.code or "", a.name.lower()))
    entities = list(db.scalars(select(LegalEntity).where(
        LegalEntity.organization_id == organization_id, LegalEntity.is_active).order_by(LegalEntity.short_name)))

    lists = wb.create_sheet("Справочники")
    lists.append(["Счёт", "Наименование счёта", "Обязательная аналитика", "", "Статья", "", "Юрлицо (ИНН)", "Юрлицо"])
    for c in "ABCEGH":
        lists[f"{c}1"].font = Font(bold=True)
    lists.column_dimensions["B"].width = 45
    lists.column_dimensions["C"].width = 45
    lists.column_dimensions["E"].width = 45
    lists.column_dimensions["G"].width = 15
    lists.column_dimensions["H"].width = 30
    for i, a in enumerate(accounts, start=2):
        lists.cell(row=i, column=1, value=a.code).number_format = "@"
        lists.cell(row=i, column=2, value=a.name)
        lists.cell(row=i, column=3, value=", ".join(ANALYTIC_TITLES[AnalyticKind(x.kind)] for x in a.analytics))
    for i, art in enumerate(articles, start=2):
        lists.cell(row=i, column=5, value=article_label(art))
    for i, e in enumerate(entities, start=2):
        lists.cell(row=i, column=7, value=e.inn).number_format = "@"
        lists.cell(row=i, column=8, value=e.short_name)

    def letter(key: str) -> str:
        return get_column_letter(next(i for i, c in enumerate(OPERATION_COLUMNS, start=1) if c.key == key))

    def add_list(key: str, ref_col: str, count: int) -> None:
        if not count:
            return
        dv = DataValidation(type="list", formula1=f"=Справочники!${ref_col}$2:${ref_col}${count + 1}",
                            allow_blank=True, showErrorMessage=False)
        ws.add_data_validation(dv)
        dv.add(f"{letter(key)}2:{letter(key)}1000")

    add_list("debit", "A", len(accounts))
    add_list("credit", "A", len(accounts))
    add_list("article", "E", len(articles))
    add_list("entity", "G", len(entities))

    _help_sheet(wb, OPERATION_COLUMNS, [
        "Загрузка операций: одна строка — одна проводка. Каждая операция загружается черновиком.",
        "Первая строка листа «Операции» — заголовки, не меняйте их. Колонки со звёздочкой обязательны.",
        "Аналитика (контрагент, договор, расчётный счёт и т.п.) ставится на ту сторону проводки, где она нужна "
        "по плану счетов (лист «Справочники»). Если аналитика нужна на обеих сторонах — ставится на обе.",
        "Строка, в которой что-то не найдено в справочниках, не загружается; номер строки и причина будут "
        "в итогах. Исправьте файл и загрузите снова — уже загруженные строки не задвоятся.",
    ])
    return _to_bytes(wb)


def counterparties_template() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Контрагенты"
    _sheet_with_header(ws, COUNTERPARTY_COLUMNS)
    _help_sheet(wb, COUNTERPARTY_COLUMNS, [
        "Загрузка контрагентов: одна строка — один контрагент.",
        "Контрагент с ИНН, который уже есть в организации, пропускается (не изменяется).",
    ])
    return _to_bytes(wb)


# --- Чтение файла ---


def _read_rows(raw: bytes, columns: list[Column]) -> list[tuple[int, dict]]:
    """Строки первого листа как {ключ колонки: значение}; колонки ищутся по заголовкам."""
    if not raw.startswith(b"PK"):
        raise ExcelError("Это не файл Excel (.xlsx). Сохраните файл в формате «Книга Excel (.xlsx)»")
    try:
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except (zipfile.BadZipFile, KeyError, ValueError, OSError) as exc:
        raise ExcelError(f"Не удалось прочитать файл Excel: {exc}") from None
    try:
        ws = wb.worksheets[0]
        rows = ws.iter_rows(values_only=True)
        header = next(rows, None)
        if header is None:
            raise ExcelError("Файл пуст")
        titles = {_norm_title(c.title): c.key for c in columns}
        positions: dict[str, int] = {}
        for i, value in enumerate(header):
            key = titles.get(_norm_title(value))
            if key is not None and key not in positions:
                positions[key] = i
        missing = [c.title for c in columns if c.required and c.key not in positions]
        if missing:
            raise ExcelError("В первой строке нет колонок: " + ", ".join(f"«{m}»" for m in missing)
                             + ". Скачайте шаблон и заполните его")
        result = []
        for number, values in enumerate(rows, start=2):
            row = {key: values[i] if i < len(values) else None for key, i in positions.items()}
            if all(_text(v) == "" for v in row.values()):
                continue
            if len(result) >= MAX_ROWS:
                raise ExcelError(f"В файле больше {MAX_ROWS} строк — разделите его на части")
            result.append((number, row))
        return result
    finally:
        wb.close()


def _norm_title(value) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip().rstrip("*").strip().lower()


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return re.sub(r"\s+", " ", str(value)).strip()


def _digits(value) -> str:
    """ИНН, номер счёта: Excel мог превратить их в число."""
    return _text(value).replace(" ", "")


def _decimal(value, title: str, places: int) -> Decimal | None:
    if value is None or _text(value) == "":
        return None
    if isinstance(value, bool):
        raise RowProblem(f"«{title}»: нужно число")
    if isinstance(value, (int, float)):
        number = Decimal(str(value))  # str(float) — кратчайшая запись, без двоичного «хвоста»
    else:
        text = _text(value).replace(" ", "").replace(" ", "").replace(",", ".")
        try:
            number = Decimal(text)
        except InvalidOperation:
            raise RowProblem(f"«{title}»: «{_text(value)}» — не число") from None
    if not number.is_finite():
        raise RowProblem(f"«{title}»: нужно число")
    quant = Decimal(1).scaleb(-places)
    rounded = round_money(number) if places == 2 else number.quantize(quant)
    if rounded != number:
        raise RowProblem(f"«{title}»: не больше {places} знаков после запятой")
    if rounded <= 0:
        raise RowProblem(f"«{title}»: должно быть больше нуля")
    if len(rounded.as_tuple().digits) > 18:
        raise RowProblem(f"«{title}»: слишком большое число")
    return rounded


def _date(value) -> dt.date:
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    text = _text(value)
    for fmt in ("%d.%m.%Y", "%d.%m.%y", "%Y-%m-%d"):
        try:
            return dt.datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    raise RowProblem(f"«Дата»: «{text}» — укажите дату в виде ДД.ММ.ГГГГ" if text else "Не указана дата")


def _account_code(value) -> str:
    if isinstance(value, float):
        # число 60.1 в ячейке — это «60.10»; 60.01 — «60.01»
        return f"{value:05.2f}" if not value.is_integer() else f"{int(value):02d}"
    if isinstance(value, int):
        return f"{value:02d}"
    text = _text(value).replace(",", ".")
    return f"0{text}" if re.fullmatch(r"\d(\.\d{2})*", text) else text


def _flag(value, title: str) -> bool:
    text = _text(value).lower()
    if text in YES:
        return True
    if text in NO:
        return False
    raise RowProblem(f"«{title}»: укажите «да» или «нет»")


# --- Операции ---


class _OperationsImporter:
    def __init__(self, db: Session, organization_id: int, user: User, branch: set[int] | None = None):
        self.db, self.org_id, self.user = db, organization_id, user
        self.branch = branch  # ветка ЦФО руководителя ЦФО: операции только с ЦФО из неё
        self.result = OperationsResult()
        self.unused: list[tuple] = []  # по каждой проводке операции: неиспользованная и использованная аналитика
        org = organization_id
        self.entities = list(db.scalars(select(LegalEntity).where(LegalEntity.organization_id == org,
                                                                  LegalEntity.is_active)))
        self.accounts = {a.code: a for a in db.scalars(select(Account))}
        self.articles = ArticleLookup(leaf_articles(db, organization_id))
        self.currencies = {c.code: c for c in db.scalars(select(Currency))}

        def scoped(model):
            return list(db.scalars(select(model).where(model.organization_id == org, model.is_active)))

        self.counterparties = scoped(Counterparty)
        self.contracts = scoped(Contract)
        self.bank_accounts = scoped(BankAccount)
        self.employees = scoped(Employee)
        self.nomenclature = scoped(Nomenclature)
        self.warehouses = scoped(Warehouse)
        self.fixed_assets = scoped(FixedAsset)
        self.cfos = cfo_lookup(db, org)
        self.executors = executor_lookup(db, org)

    # --- поиск в справочниках ---

    @staticmethod
    def _one(items: list, what: str, value: str):
        if not items:
            raise RowProblem(f"{what} «{value}» не найден(а) в справочнике организации")
        if len(items) > 1:
            raise RowProblem(f"{what} «{value}»: найдено несколько — укажите точнее (ИНН, номер, артикул)")
        return items[0]

    @staticmethod
    def _by_name(items: list, value: str, attr: str = "name") -> list:
        low = value.lower()
        return [x for x in items if getattr(x, attr).strip().lower() == low]

    def entity(self, value) -> LegalEntity:
        text = _digits(value)
        if not text:
            if len(self.entities) == 1:
                return self.entities[0]
            raise RowProblem("Не указано юрлицо (ИНН) — в организации их несколько")
        found = [e for e in self.entities if e.inn == text] or self._by_name(self.entities, _text(value), "short_name")
        return self._one(found, "Юрлицо", _text(value))

    def account(self, value, title: str) -> Account:
        if _text(value) == "":
            raise RowProblem(f"Не указан «{title}»")
        code = _account_code(value)
        account = self.accounts.get(code)
        if account is None:
            raise RowProblem(f"«{title}»: счёта {code} нет в плане счетов")
        return account

    def resolve(self, kind: str, value, entity: LegalEntity, counterparty: Counterparty | None) -> int:
        text = _text(value)
        if kind == "counterparty":
            digits = _digits(value)
            found = [c for c in self.counterparties if c.inn == digits] if re.fullmatch(r"\d{10}|\d{12}", digits) else []
            return self._one(found or self._by_name(self.counterparties, text), "Контрагент", text).id
        if kind == "contract":
            if counterparty is None:
                raise RowProblem("Договор указан без контрагента")
            found = [c for c in self.contracts if c.counterparty_id == counterparty.id
                     and c.legal_entity_id == entity.id and c.number.strip().lower() == text.lower()]
            if not found:
                raise RowProblem(f"Договор № {text} с «{counterparty.name}» у этого юрлица не найден")
            return self._one(found, "Договор №", text).id
        if kind == "bank_account":
            digits = _digits(value)
            return self._one([b for b in self.bank_accounts if b.number == digits], "Расчётный счёт", text).id
        if kind == "employee":
            found = [e for e in self.employees if e.personnel_number and e.personnel_number == text]
            return self._one(found or self._by_name(self.employees, text, "full_name"), "Сотрудник", text).id
        if kind == "nomenclature":
            found = [n for n in self.nomenclature if n.article and n.article.strip().lower() == text.lower()]
            return self._one(found or self._by_name(self.nomenclature, text), "Номенклатура", text).id
        if kind == "warehouse":
            return self._one(self._by_name(self.warehouses, text), "Склад", text).id
        if kind == "fixed_asset":
            found = [a for a in self.fixed_assets if a.inventory_number and a.inventory_number == text]
            return self._one(found or self._by_name(self.fixed_assets, text), "Основное средство", text).id
        if kind == "cfo":
            return self.cfos.find(text)
        if kind == "executor":
            return self.executors.find(text)
        raise AssertionError(kind)

    # --- строка → проводка ---

    def posting(self, row: dict, entity: LegalEntity, article: Article | None) -> PostingIn:
        debit = self.account(row.get("debit"), "Счёт Дт")
        credit = self.account(row.get("credit"), "Счёт Кт")
        amount = _decimal(row.get("amount"), "Сумма", 2)
        if amount is None:
            raise RowProblem("Не указана сумма")
        quantity = _decimal(row.get("quantity"), "Количество", 3)
        code = _text(row.get("currency")).upper()
        currency = None
        if code:
            currency = self.currencies.get(code)
            if currency is None:
                raise RowProblem(f"«Валюта»: валюты {code} нет в справочнике")

        values: dict[str, int] = {}
        counterparty = None
        for kind in ANALYTIC_KEYS:
            if _text(row.get(kind)) == "":
                continue
            values[kind] = self.resolve(kind, row.get(kind), entity, counterparty)
            if kind == "counterparty":
                counterparty = next(c for c in self.counterparties if c.id == values[kind])

        sides = {"debit": AnalyticsIn(), "credit": AnalyticsIn()}
        used: set[str] = set()
        for side, account in (("debit", debit), ("credit", credit)):
            required = {a.kind for a in account.analytics}
            required |= {extra for extra, base in OPTIONAL_ANALYTICS.items() if base in required}
            for kind, ref_id in values.items():
                if kind in required:
                    setattr(sides[side], kind, ref_id)
                    used.add(kind)
            if article is not None:
                if "cash_flow_item" in required and article.is_cash_flow:
                    sides[side].cash_flow_item = article.id
                if "pl_item" in required and article.is_pl:
                    sides[side].pl_item = article.id
        self.unused.append(([k for k in values if k not in used], used, debit.code, credit.code))
        return PostingIn(
            debit_account_id=debit.id, credit_account_id=credit.id, amount=amount,
            currency_id=currency.id if currency else None, quantity=quantity,
            debit=sides["debit"], credit=sides["credit"],
        )

    def article(self, value) -> Article | None:
        text = _text(value)
        if not text:
            return None
        return self.articles.find(text, "Статья")

    # --- группа строк → черновик ---

    def import_group(self, rows: list[tuple[int, dict]], occurrence: int) -> None:
        first_no, first = rows[0]
        try:
            op_date = _date(first.get("date"))
            entity = self.entity(first.get("entity"))
            article = self.article(first.get("article"))
            postings = []
            self.unused = []
            for number, row in rows:
                try:
                    if _date(row.get("date")) != op_date or self.entity(row.get("entity")).id != entity.id:
                        raise RowProblem("В одной операции дата и юрлицо должны совпадать во всех строках")
                    postings.append(self.posting(row, entity, self.article(row.get("article")) or article))
                except RowProblem as exc:
                    if len(rows) == 1:
                        raise
                    raise RowProblem(f"строка {number}: {exc}") from None
            if len(postings) > 100:
                raise RowProblem("В одной операции больше 100 проводок")
            self.check_unused(rows)
            if self.branch is not None:
                cfos = {c for p in postings for c in (p.debit.cfo, p.credit.cfo) if c is not None}
                if not cfos & self.branch or cfos - self.branch:
                    raise RowProblem("Можно загружать только операции с ЦФО вашей ветки")
        except RowProblem as exc:
            self.result.errors.append(RowError(first_no, str(exc)))
            return

        key = self.import_key(rows, occurrence)
        if self.db.scalar(select(Operation.id).where(Operation.organization_id == self.org_id,
                                                     Operation.import_key == key)):
            self.result.skipped_duplicates += 1
            return
        description = "; ".join(dict.fromkeys(_text(r.get("description")) for _, r in rows if _text(r.get("description"))))
        body = OperationIn(legal_entity_id=entity.id, op_date=op_date, article_id=article.id if article else None,
                           description=description[:2000], postings=postings)
        op = Operation(organization_id=self.org_id, created_by=self.user.id, status="draft", import_key=key)
        try:
            journal.save_draft(self.db, op, body, self.user)
        except JournalError as exc:
            self.result.errors.append(RowError(first_no, str(exc)))
            return
        op.number = journal.next_number(self.db, self.org_id)
        self.db.add(op)
        self.db.flush()
        journal._event(self.db, op, self.user, "created")
        self.result.created.append(op.number)
        if journal.posting_errors(self.db, op):
            self.result.incomplete.append(op.number)

    def check_unused(self, rows: list[tuple[int, dict]]) -> None:
        """Аналитика из файла, которая не легла ни на одну проводку операции, — ошибка, а не тихая потеря.
        В строках одной операции аналитику можно повторять: она ставится только туда, где нужна."""
        used_anywhere = set().union(*(u[1] for u in self.unused))
        for (number, _), (unused, _, debit, credit) in zip(rows, self.unused):
            lost = [k for k in unused if k not in used_anywhere]
            if lost:
                titles = ", ".join(f"«{ANALYTIC_TITLES[AnalyticKind(k)]}»" for k in lost)
                text = f"указано {titles}, но на счетах {debit} и {credit} эта аналитика не ведётся"
                raise RowProblem(f"строка {number}: {text}" if len(rows) > 1 else text[0].upper() + text[1:])

    @staticmethod
    def import_key(rows: list[tuple[int, dict]], occurrence: int) -> str:
        """Содержимое строк (без номеров строк) + номер повторения такой же операции в файле."""
        content = "\n".join("|".join(f"{k}={_text(v)}" for k, v in sorted(r.items())) for _, r in rows)
        digest = hashlib.sha256(content.encode()).hexdigest()[:40]
        return f"xlsx:{digest}:{occurrence}"


def import_operations(db: Session, organization_id: int, user: User, raw: bytes,
                      branch: set[int] | None = None) -> OperationsResult:
    rows = _read_rows(raw, OPERATION_COLUMNS)
    importer = _OperationsImporter(db, organization_id, user, branch)
    importer.result.rows = len(rows)

    groups: dict[str, list] = {}
    order: list[list] = []
    for number, row in rows:
        group = _text(row.get("group"))
        if not group:
            order.append([(number, row)])
        elif group in groups:
            groups[group].append((number, row))
        else:
            groups[group] = [(number, row)]
            order.append(groups[group])

    seen: dict[str, int] = {}
    for group_rows in order:
        base = importer.import_key(group_rows, 0)
        seen[base] = seen.get(base, -1) + 1
        importer.import_group(group_rows, seen[base])
    return importer.result


# --- Контрагенты ---


def import_counterparties(db: Session, organization_id: int, raw: bytes) -> CounterpartiesResult:
    rows = _read_rows(raw, COUNTERPARTY_COLUMNS)
    result = CounterpartiesResult(rows=len(rows))
    existing = list(db.scalars(select(Counterparty).where(Counterparty.organization_id == organization_id)))
    by_inn = {c.inn for c in existing if c.inn}
    names_without_inn = {c.name.strip().lower() for c in existing if not c.inn}

    for number, row in rows:
        try:
            name = _text(row.get("name"))[:255]
            if not name:
                raise RowProblem("Не указано наименование")
            inn = _digits(row.get("inn")) or None
            if inn is not None and not re.fullmatch(r"\d{10}|\d{12}", inn):
                raise RowProblem(f"«ИНН»: «{_text(row.get('inn'))}» — нужно 10 или 12 цифр")
            kpp = _digits(row.get("kpp")) or None
            if kpp is not None and not re.fullmatch(r"\d{9}", kpp):
                raise RowProblem("«КПП»: нужно 9 цифр")
            kind_text = _text(row.get("kind")).lower()
            if kind_text:
                kind = KIND_TITLES.get(kind_text)
                if kind is None:
                    raise RowProblem("«Вид»: укажите «Организация», «ИП» или «Физлицо»")
            elif inn is None:
                kind = "person"
            else:
                kind = "legal" if len(inn) == 10 else "entrepreneur"
            try:
                check_inn_for_kind(kind, inn)
            except ValueError as exc:
                raise RowProblem(f"«ИНН»: {exc}") from None
            is_customer = _flag(row.get("is_customer"), "Покупатель")
            is_supplier = _flag(row.get("is_supplier"), "Поставщик")
        except RowProblem as exc:
            result.errors.append(RowError(number, str(exc)))
            continue

        if (inn and inn in by_inn) or (not inn and name.lower() in names_without_inn):
            result.skipped_existing += 1
            continue
        db.add(Counterparty(
            organization_id=organization_id, name=name, full_name=_text(row.get("full_name"))[:500],
            kind=kind, inn=inn, kpp=kpp, is_customer=is_customer, is_supplier=is_supplier,
        ))
        if inn:
            by_inn.add(inn)
        else:
            names_without_inn.add(name.lower())
        result.created.append(name)
    db.flush()
    return result
