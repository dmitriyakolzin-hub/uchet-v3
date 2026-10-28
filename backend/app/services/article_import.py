"""Загрузка справочника статей из Excel с совмещением (docs/TZ.md, раздел 19, «Детали 9.2»).

Сначала `preview` — разбор файла без изменений (вид строки, родитель по коду, предлагаемое совмещение), затем `apply`
со строками, которые пользователь проверил (вид и совмещение можно поменять).
"""

import io
import re
import zipfile
from dataclasses import dataclass, field

from openpyxl import load_workbook
from sqlalchemy import delete, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.models import Account, Article, Posting, ReportFormRow
from app.services import articles as tree
from app.services.excel_import import _text

# Показатели (расчётные строки и остатки) по коду: точные коды и разделы целиком
MAX_ROWS = 2000
INDICATOR_CODES = {"1.0", "2.0", "4.0", "4.1", "4.1.1", "5.0", "5.1", "5.2", "5.3"}
INDICATOR_SECTIONS = ("3.", "6.", "5.0.", "5.2.", "5.3.")

OP, INV, FIN = "operating", "investing", "financing"
IN, OUT = "inflow", "outflow"


@dataclass
class Rule:
    flow: tuple[str, str] | None  # (вид деятельности, направление) или None — не ДДС
    section: str | None  # раздел ОПиУ или None — не ОПиУ
    debit: str | None = None
    credit: str | None = None


# Признаки по разделу кода: побеждает самый длинный подходящий префикс
RULES: dict[str, Rule] = {
    "1.1": Rule((OP, IN), "revenue", "62.01", "90.01"),
    "1.2": Rule((OP, IN), "revenue", "62.01", "90.01"),
    "1.3": Rule((OP, IN), "other_income", "76", "91.01"),
    "1": Rule((OP, IN), "revenue", "62.01", "90.01"),
    "2.1": Rule(None, "cost", "90.02", "41"),
    "2.2": Rule(None, "selling", "44", "70"),
    "2.2.2": Rule(None, "selling", "44", "69"),
    "2.9": Rule((FIN, OUT), None, "75", "51"),
    "2.10": Rule((FIN, OUT), None, "75", "51"),
    "2.11": Rule((FIN, OUT), None, "75", "51"),
    "2.12": Rule(None, "income_tax", "99", "68.04"),
    "2": Rule((OP, OUT), "selling", "44", "60.01"),
    "4.2": Rule((FIN, OUT), None, "76", "51"),
    "4.3": Rule((INV, OUT), None, "08", "60.01"),
    "5.1.1": Rule((FIN, IN), None, "51", "66"),
    "5.1.2": Rule((FIN, OUT), None, "66", "51"),
}
GOODS_RULE = Rule(None, "selling", "44", "41")  # пометка «товары»: расход товарами, без денег

# Соответствия названий файла и стандартного набора (нижний регистр)
SYNONYMS = {
    "выручка без скидок": "выручка от продаж",
    "общая сырьевая себестоимость": "себестоимость продаж",
    "страховые взносы": "страховые взносы с оплаты труда",
    "коммунальные расходы": "коммунальные услуги",
    "комиссия банка": "банковские комиссии",
    "реклама": "реклама и маркетинг",
    "связь": "связь и it",
    "оборудование": "инвестиции: оборудование",
    "строительство": "инвестиции: ремонт и строительство",
    "прочие капиталовложения": "инвестиции: прочие вложения",
    "поступление": "получение кредитов и займов",
    "погашение": "погашение кредитов и займов",
}


# Стартовые формулы показателей (решения владельца, docs/TZ.md, раздел 19) и строки с обратным знаком
DEFAULT_FORMULAS = {
    "1.0": "1.1 - 1.2",
    "2.0": "2.2 + 2.3 + 2.4 + 2.5 + 2.6 + 2.7 + 2.8",
    "3.1": "1.0 - 2.1",
    "3.2": "3.1 - 2.0",
    "3.3": "3.2 + 1.3",
    "3.4": "3.3 - 2.12",
    "4.1.1": "3.4",
    "4.1": "4.1.1",
    "3.5": "4.1 - 4.2",
    "3.6": "3.5 - 4.3",
}
DEFAULT_INVERT = {"1.2"}


class ArticleImportError(Exception):
    pass


@dataclass
class Row:
    row: int
    code: str | None
    name: str
    note: str = ""
    kind: str = "article"  # article, indicator, header
    parent_code: str | None = None
    match_id: int | None = None
    match_name: str = ""
    match_reason: str = ""  # «по коду», «по названию», «по соответствию», «вручную»
    warning: str = ""
    flags: str = ""  # признаки по правилу, текстом


@dataclass
class ApplyResult:
    created: int = 0
    merged: int = 0
    updated: int = 0
    skipped: int = 0
    errors: list[tuple[int, str]] = field(default_factory=list)


def norm_code(value) -> str | None:
    text = re.sub(r"\s+", "", _text(value)).rstrip(".")
    return text or None


def norm_name(value: str) -> str:
    return re.sub(r"\s+", " ", value or "").strip().lower()


def kind_of(code: str | None) -> str:
    if code is None:
        return "header"
    if code in INDICATOR_CODES or any((code + ".").startswith(s) for s in INDICATOR_SECTIONS):
        return "indicator"
    return "article"


def rule_for(code: str, note: str = "") -> Rule | None:
    if note.strip().lower() == "товары":
        return GOODS_RULE
    best = None
    for prefix, rule in RULES.items():
        if code == prefix or code.startswith(prefix + "."):
            if best is None or len(prefix) > len(best[0]):
                best = (prefix, rule)
    return best[1] if best else None


def describe(rule: Rule | None) -> str:
    if rule is None:
        return "признаки — вручную"
    parts = []
    if rule.flow:
        parts.append("ДДС " + {"operating": "опер.", "investing": "инвест.", "financing": "фин."}[rule.flow[0]]
                     + (" поступление" if rule.flow[1] == IN else " выплата"))
    if rule.section:
        parts.append("ОПиУ")
    if rule.debit:
        parts.append(f"Дт {rule.debit} — Кт {rule.credit}")
    return ", ".join(parts)


def _sheet_rows(raw: bytes) -> list[list]:
    """Значения первого листа построчно."""
    if not raw.startswith(b"PK"):
        raise ArticleImportError("Это не файл Excel (.xlsx). Сохраните файл в формате «Книга Excel (.xlsx)»")
    try:
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except (zipfile.BadZipFile, KeyError, ValueError, OSError) as exc:
        raise ArticleImportError(f"Не удалось прочитать файл Excel: {exc}") from None
    try:
        rows = [list(values) for _, values in zip(range(MAX_ROWS), wb.worksheets[0].iter_rows(values_only=True))]
    finally:
        wb.close()
    return rows


def parse(raw: bytes) -> list[Row]:
    sheet = _sheet_rows(raw)
    header_row, name_col, code_col = None, None, None
    for r, values in enumerate(sheet[:20], start=1):
        titles = [_text(v).lower() for v in values]
        names = [i for i, t in enumerate(titles) if t.startswith("наименование")]
        codes = [i for i, t in enumerate(titles) if t.startswith("код") or t.startswith("статья")]
        if names and codes:
            header_row, name_col, code_col = r, names[0], codes[0]
            break
    if header_row is None:
        raise ArticleImportError("Не найдена строка заголовков: нужны колонки «Наименование…» и «Код» (или «Статья…»)")
    note_col = max(name_col, code_col) + 1
    rows = []
    for r, values in enumerate(sheet[header_row:], start=header_row + 1):
        values = values + [None] * 3
        name = re.sub(r"\s+", " ", _text(values[name_col])).strip()
        if not name:
            continue
        code = norm_code(values[code_col])
        rows.append(Row(row=r, code=code, name=name[:255], note=_text(values[note_col])[:100], kind=kind_of(code)))
    if not rows:
        raise ArticleImportError("В файле нет строк со статьями")
    return rows


def _parent_codes(rows: list[Row]) -> None:
    articles = {r.code for r in rows if r.kind == "article"}
    for r in rows:
        r.parent_code = None
        if r.kind == "article" and r.code and "." in r.code:
            parent = r.code.rsplit(".", 1)[0]
            r.parent_code = parent if parent in articles else None


def _has_children(rows: list[Row], code: str) -> bool:
    return any(r.parent_code == code for r in rows)


def _used_flags(db: Session, article_id: int) -> tuple[bool, bool]:
    """Есть ли у статьи проводки как у статьи ДДС / как у статьи ОПиУ."""
    cash = db.scalar(select(Posting.id).where((Posting.dt_cash_flow_article_id == article_id)
                                              | (Posting.ct_cash_flow_article_id == article_id)).limit(1))
    pl = db.scalar(select(Posting.id).where((Posting.dt_pl_article_id == article_id)
                                            | (Posting.ct_pl_article_id == article_id)).limit(1))
    return cash is not None, pl is not None


def merge_problem(db: Session, rows: list[Row], row: Row, article: Article) -> str:
    """Почему строку нельзя совместить со статьёй (пусто — можно)."""
    if article.system_key and _has_children(rows, row.code):
        return "системная статья не может стать группой"
    if article.is_transfer:
        return "статья перевода между своими счетами не совмещается"
    rule = _effective_rule(rows, row)
    if rule is None:
        return ""
    used_cash, used_pl = _used_flags(db, article.id)
    if (used_cash and rule.flow is None) or (used_pl and rule.section is None):
        return "признаки не совпадают: по статье есть проводки как по статье " + ("ДДС" if used_cash and rule.flow is None
                                                                                   else "ОПиУ")
    if rule.flow and article.is_cash_flow and article.cash_flow_direction != rule.flow[1] and used_cash:
        return "признаки не совпадают: другое направление ДДС"
    return ""


def _effective_rule(rows: list[Row], row: Row) -> Rule | None:
    """Правило строки: подстатья берёт признаки верхней статьи своей ветки (пометка «товары» — только шаблон)."""
    by_code = {r.code: r for r in rows}
    top = row
    while top.parent_code and top.parent_code in by_code:
        top = by_code[top.parent_code]
    return rule_for(top.code or "", top.note)


def preview(db: Session, organization_id: int, raw: bytes) -> list[Row]:
    rows = parse(raw)
    _parent_codes(rows)
    existing = list(db.scalars(select(Article).where(Article.organization_id == organization_id)))
    by_code = {a.code.rstrip("."): a for a in existing if a.code}
    by_name: dict[str, list[Article]] = {}
    for a in existing:
        by_name.setdefault(norm_name(a.name), []).append(a)
    taken: set[int] = set()
    for r in rows:
        if r.kind != "article":
            continue
        r.flags = describe(_effective_rule(rows, r))
        candidates = []
        if r.code in by_code:
            candidates.append((by_code[r.code], "по коду"))
        for a in by_name.get(norm_name(r.name), []):
            candidates.append((a, "по названию"))
        synonym = SYNONYMS.get(norm_name(r.name))
        for a in by_name.get(synonym, []) if synonym else []:
            candidates.append((a, "по соответствию"))
        for article, reason in candidates:
            if article.id in taken:
                continue
            if reason != "по коду" and article.code and article.code.rstrip(".") != r.code:
                continue  # у статьи уже другой код — это другая статья
            problem = merge_problem(db, rows, r, article)
            if problem:
                r.warning = f"не совмещено с «{article.name}»: {problem}"
                continue
            r.match_id, r.match_name, r.match_reason = article.id, article.name, reason
            taken.add(article.id)
            break
    return rows


def _template(db: Session, rule: Rule | None) -> tuple[int | None, int | None]:
    if rule is None or rule.debit is None:
        return None, None
    ids = dict(db.execute(select(Account.code, Account.id).where(Account.code.in_((rule.debit, rule.credit)))).all())
    return ids.get(rule.debit), ids.get(rule.credit)


def apply(db: Session, organization_id: int, rows: list[Row]) -> ApplyResult:
    """Статьи из проверенных строк: родители раньше подстатей; совмещённые — обновляются, остальные — создаются."""
    result = ApplyResult()
    _parent_codes(rows)
    articles_rows = [r for r in rows if r.kind == "article"]
    result.skipped = len(rows) - len(articles_rows)
    seen_codes = set()
    for r in articles_rows:
        if not r.code:
            result.errors.append((r.row, "у статьи нет кода"))
        elif r.code in seen_codes:
            result.errors.append((r.row, f"код {r.code} повторяется в файле"))
        seen_codes.add(r.code)
    if result.errors:
        return result
    depth = {r.code: r.code.count(".") for r in articles_rows}
    created: dict[str, Article] = {}
    order = {r.code: i for i, r in enumerate(rows, start=1)}
    for r in sorted(articles_rows, key=lambda x: (depth[x.code], order[x.code])):
        parent = created.get(r.parent_code) if r.parent_code else None
        if r.parent_code and parent is None:
            result.errors.append((r.row, f"не загружена родительская статья {r.parent_code}"))
            continue
        try:
            with db.begin_nested():
                if r.match_id is not None:
                    article = db.get(Article, r.match_id)
                    if article is None or article.organization_id != organization_id:
                        raise ArticleImportError("статья для совмещения не найдена")
                    problem = merge_problem(db, rows, r, article)
                    if problem:
                        raise ArticleImportError(problem)
                    creating = False
                    outcome = "updated" if article.code and article.code.rstrip(".") == r.code else "merged"
                else:
                    article = Article(organization_id=organization_id)
                    db.add(article)
                    creating, outcome = True, "created"
                with db.no_autoflush:  # статья пишется в базу только целиком
                    article.name, article.code, article.sort_order = r.name, r.code + ".", order[r.code]
                    article.parent_id = parent.id if parent else None
                    if parent is None:
                        rule = rule_for(r.code, r.note)
                        if rule is not None:
                            article.is_cash_flow = rule.flow is not None
                            article.cash_flow_activity, article.cash_flow_direction = rule.flow or (None, None)
                            article.is_pl = rule.section is not None
                            article.pl_section = rule.section
                            if creating or article.debit_account_id is None:
                                article.debit_account_id, article.credit_account_id = _template(db, rule)
                        elif creating:
                            raise ArticleImportError("для этого кода нет правила признаков — создайте статью вручную")
                    if parent is not None and r.note.strip().lower() == "товары" and creating:
                        article.debit_account_id, article.credit_account_id = _template(db, GOODS_RULE)
                    tree.check_tree(db, article, creating)
                db.flush()
                created[r.code] = article
            setattr(result, outcome, getattr(result, outcome) + 1)
        except (ArticleImportError, tree.ArticleError) as exc:
            result.errors.append((r.row, str(exc)))
        except IntegrityError:
            result.errors.append((r.row, "статья с таким кодом или таким названием в этой группе уже есть"))
    save_form(db, organization_id, rows, created)
    return result


def save_form(db: Session, organization_id: int, rows: list[Row], created: dict[str, Article]) -> None:
    """Форма управленческого отчёта — все строки файла по порядку; исправленные формулы и знак сохраняются."""
    old = {r.code: r for r in db.scalars(select(ReportFormRow).where(ReportFormRow.organization_id == organization_id))
           if r.code}
    db.execute(delete(ReportFormRow).where(ReportFormRow.organization_id == organization_id))
    for position, r in enumerate(rows, start=1):
        kept = old.get(r.code)
        article = created.get(r.code) if r.kind == "article" else None
        db.add(ReportFormRow(
            organization_id=organization_id, position=position, code=r.code, name=r.name, kind=r.kind,
            article_id=article.id if article else None, note=r.note,
            formula=kept.formula if kept and kept.kind == r.kind else
            (DEFAULT_FORMULAS.get(r.code, "") if r.kind == "indicator" else ""),
            invert=kept.invert if kept else r.code in DEFAULT_INVERT,
        ))
    db.flush()

