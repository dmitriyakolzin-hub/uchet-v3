"""Начисление зарплаты: дни, НДФЛ по шкале, взносы, проведение в журнал (docs/TZ.md, раздел 14)."""

import datetime as dt
from decimal import Decimal

import pytest

from app.core import payroll_rates as rates


def test_rates():
    assert rates.ndfl_total(Decimal("120000")) == Decimal("15600")
    assert rates.ndfl_total(Decimal("3000000")) == Decimal("402000")  # 13% до 2,4 млн + 15% сверху
    assert rates.ndfl_total(Decimal("60000000")) == Decimal(
        "312000") + Decimal("390000") + Decimal("2700000") + Decimal("6000000") + Decimal("2200000")
    assert rates.insurance("sme", 2026, Decimal("120000"), Decimal(0)) == Decimal("24095.93")
    assert rates.insurance("sme", 2026, Decimal("30000"), Decimal(0)) == Decimal("9000.00")
    assert rates.insurance("general", 2026, Decimal("100000"), Decimal("2950000")) == Decimal("19421.00")
    assert rates.mrot(2030) == rates.MROT[2026]  # для будущих лет — последние известные
    assert rates.working_days(dt.date(2026, 2, 1), dt.date(2026, 2, 28)) == 19  # 23 февраля — праздник


@pytest.fixture
def env(client, headers_for, make_user, make_org, db):
    from app.models import LegalEntity

    fd = make_user(login="fd")
    ceo = make_user(login="ceo")
    other = make_user(login="other")
    org = make_org("Группа А", members=[(fd, "finance_director"), (ceo, "ceo")])
    make_org("Клиент Б", members=[(other, "finance_director")])
    le = LegalEntity(organization_id=org.id, full_name="ООО А", short_name="А", inn="7701234567")
    le2 = LegalEntity(organization_id=org.id, full_name="ООО Б", short_name="Б", inn="7701234568")
    db.add_all([le, le2])
    db.commit()
    h = headers_for(fd)
    o = f"/api/organizations/{org.id}"
    art = {a["name"]: a["id"] for a in client.get(f"/api/organizations/{org.id}/articles", headers=h).json()}
    shop = client.post(f"{o}/cfos", headers=h, json={"name": "Магазин"}).json()
    sales = art["Оплата труда — коммерческий персонал"]

    def employee(name, salary, **fields):
        emp = client.post(f"{o}/employees", headers=h, json={"full_name": name, "legal_entity_id": le.id,
                                                            "cfo_id": shop["id"], "pl_item_id": sales, **fields})
        assert emp.status_code == 201, emp.text
        emp = emp.json()
        client.post(f"/api/employees/{emp['id']}/salaries", headers=h, json={"valid_from": "2026-01-01", "amount": salary})
        return emp

    a = employee("Андреев А.", "100000")
    b = employee("Борисов Б.", "60000", hired_on="2026-02-16")
    employee("Ведёт другое юрлицо", "50000", legal_entity_id=le2.id)
    employee("Уволен в январе", "50000", fired_on="2026-01-31")
    acc = {x["code"]: x["id"] for x in client.get("/api/accounts", headers=h).json()}
    return {"o": o, "h": h, "le": le, "le2": le2, "a": a, "b": b, "shop": shop, "art": art, "acc": acc,
            "ceo": headers_for(ceo), "other": headers_for(other)}


def create(client, env, month=2):
    resp = client.post(f"{env['o']}/payroll-runs", headers=env["h"],
                       json={"legal_entity_id": env["le"].id, "year": 2026, "month": month})
    assert resp.status_code == 201, resp.text
    return resp.json()


def line_of(run, emp):
    return next(line for line in run["lines"] if line["employee_id"] == emp["id"])


def test_fill_and_calculate(client, env):
    run = create(client, env)
    assert run["norm_days"] == 19 and run["insurance_tariff"] == "sme"
    assert sorted(line["employee_name"] for line in run["lines"]) == ["Андреев А.", "Борисов Б."]
    b = line_of(run, env["b"])  # принят 16 февраля: 9 рабочих дней из 19
    assert (b["work_days"], b["salary_amount"]) == ("9.00", "28421.05")

    a = line_of(run, env["a"])
    run = client.patch(f"/api/payroll-lines/{a['id']}", headers=env["h"], json={"bonus": "20000"}).json()
    a = line_of(run, env["a"])
    assert (a["accrued"], a["ndfl"], a["insurance"], a["net"]) == ("120000.00", "15600.00", "24095.93", "104400.00")
    assert run["totals"]["accrued"] == "148421.05"
    assert create_conflict(client, env) == 409

    # неполный месяц и удержание
    run = client.patch(f"/api/payroll-lines/{a['id']}", headers=env["h"],
                       json={"work_days": "9.5", "other_deduction": "1000"}).json()
    a = line_of(run, env["a"])
    assert a["salary_amount"] == "50000.00" and a["net"] == str(Decimal(a["accrued"]) - Decimal(a["ndfl"]) - 1000)
    # поправили норму: у кого была полная норма — тоже меняется
    run = client.patch(f"/api/payroll-runs/{run['id']}", headers=env["h"], json={"norm_days": 18}).json()
    assert line_of(run, env["b"])["work_days"] == "9.00"


def create_conflict(client, env):
    return client.post(f"{env['o']}/payroll-runs", headers=env["h"],
                       json={"legal_entity_id": env["le"].id, "year": 2026, "month": 2}).status_code


def test_post_creates_journal_operation(client, env, db):
    from app.models import Employee

    run = create(client, env)
    # сотрудник без ЦФО не даёт провести
    db.query(Employee).filter_by(id=env["b"]["id"]).update({"cfo_id": None})
    db.commit()
    resp = client.post(f"/api/payroll-runs/{run['id']}/post", headers=env["h"])
    assert resp.status_code == 422 and "Борисов Б." in resp.json()["detail"]
    db.query(Employee).filter_by(id=env["b"]["id"]).update({"cfo_id": env["shop"]["id"]})
    db.commit()

    posted = client.post(f"/api/payroll-runs/{run['id']}/post", headers=env["h"])
    assert posted.status_code == 200, posted.text
    posted = posted.json()
    assert posted["status"] == "posted"
    op = client.get(f"/api/operations/{posted['operation_id']}", headers=env["h"]).json()
    assert op["status"] == "posted" and op["op_date"] == "2026-02-28"
    codes = {v: k for k, v in env["acc"].items()}
    kinds = {(codes[p["debit_account_id"]], codes[p["credit_account_id"]]) for p in op["postings"]}
    assert kinds == {("44", "70"), ("70", "68.01"), ("44", "69")}
    # проведённое не меняется
    a = line_of(posted, env["a"])
    assert client.patch(f"/api/payroll-lines/{a['id']}", headers=env["h"], json={"bonus": "1"}).status_code == 422

    # ОПиУ: оплата труда — по статье сотрудника (коммерческие), взносы — по статье взносов (управленческие)
    pl = client.get(f"{env['o']}/reports/pl", headers=env["h"],
                    params={"date_from": "2026-02-01", "date_to": "2026-02-28"}).json()
    rows = {r["key"]: r for r in pl["rows"]}
    assert Decimal(rows["selling"]["values"]["total"]) == -Decimal(posted["totals"]["accrued"])
    assert Decimal(rows["admin"]["values"]["total"]) == -Decimal(posted["totals"]["insurance"])


def test_ndfl_cumulative_and_unpost_order(client, env):
    feb = create(client, env)
    a = line_of(feb, env["a"])
    client.patch(f"/api/payroll-lines/{a['id']}", headers=env["h"], json={"bonus": "20000"})
    assert client.post(f"/api/payroll-runs/{feb['id']}/post", headers=env["h"]).status_code == 200
    mar = create(client, env, month=3)
    a = line_of(mar, env["a"])
    assert a["ndfl"] == "13000.00"  # ndfl(220 000) − уже удержано 15 600
    assert client.post(f"/api/payroll-runs/{mar['id']}/post", headers=env["h"]).status_code == 200
    resp = client.post(f"/api/payroll-runs/{feb['id']}/unpost", headers=env["h"])
    assert resp.status_code == 422 and "следующих месяцев" in resp.json()["detail"]
    assert client.post(f"/api/payroll-runs/{mar['id']}/unpost", headers=env["h"]).json()["status"] == "draft"
    unposted = client.post(f"/api/payroll-runs/{feb['id']}/unpost", headers=env["h"]).json()
    assert unposted["status"] == "draft" and unposted["operation_id"] is None
    assert client.get(f"{env['o']}/operations", headers=env["h"]).json()["items"] == []
    assert client.delete(f"/api/payroll-runs/{feb['id']}", headers=env["h"]).status_code == 204


def test_general_tariff_and_access(client, env):
    client.patch(f"/api/legal-entities/{env['le'].id}", headers=env["h"], json={"insurance_tariff": "general"})
    run = create(client, env)
    assert run["insurance_tariff"] == "general"
    assert line_of(run, env["a"])["insurance"] == "30000.00"
    assert client.get(f"{env['o']}/payroll-runs", headers=env["ceo"]).status_code == 403
    assert client.get(f"/api/payroll-runs/{run['id']}", headers=env["ceo"]).status_code == 403
    assert client.get(f"/api/payroll-runs/{run['id']}", headers=env["other"]).status_code == 404
    listed = client.get(f"{env['o']}/payroll-runs", headers=env["h"], params={"year": 2026}).json()
    assert [(r["month"], r["totals"]["accrued"]) for r in listed] == [(2, run["totals"]["accrued"])]


def test_lines_excel_roundtrip(client, env):
    import io

    from openpyxl import load_workbook

    run = create(client, env)
    resp = client.get(f"/api/payroll-runs/{run['id']}/lines.xlsx", headers=env["h"])
    wb = load_workbook(io.BytesIO(resp.content))
    ws = wb.active
    assert [c.value for c in ws[1]][:4] == ["Таб. №", "Вн. №", "Сотрудник", "Отработано дней"]
    for row in ws.iter_rows(min_row=2):
        if row[2].value == "Андреев А.":
            row[4].value = "20 000,50"  # премия
            row[8].value = 500  # удержание
    ws.append([None, None, "Нет такого", 5])
    ws.append([None, None, "Уволен в январе", 0])  # сотрудник юрлица, не попавший в месяц, — добавляется
    out = io.BytesIO()
    wb.save(out)
    result = client.post(f"/api/payroll-runs/{run['id']}/lines", headers={**env["h"], "Content-Type":
                         "application/octet-stream"}, content=out.getvalue()).json()
    assert result["changed"] == 3 and [e["row"] for e in result["errors"]] == [4]
    a = line_of(result["run"], env["a"])
    assert (a["bonus"], a["other_deduction"], a["accrued"]) == ("20000.50", "500.00", "120000.50")
    assert len(result["run"]["lines"]) == 3


def test_lines_ready_amounts(client, env, db):
    """Готовые суммы из другой программы: есть колонка «НДФЛ» — не пересчитываются."""
    from app.models import Employee

    from tests.test_excel_import import xlsx

    db.query(Employee).filter_by(id=env["b"]["id"]).one().internal_number = "M-2"
    db.commit()
    run = create(client, env)
    header = ["Вн. №", "Сотрудник", "Премия", "Начислено", "НДФЛ", "Взносы", "Прочие удержания", "К выплате"]
    raw = xlsx(header, [
        [None, "Андреев А.", 10000, 110000, 14300, None, 700, 95000],  # взносы посчитает программа
        ["M-2", None, None, "30 000", 3900, 9000, None, None],
        [None, "Андреев А.", None, 1000, 130, None, None, 999],  # «к выплате» не сходится
    ])
    resp = client.post(f"/api/payroll-runs/{run['id']}/lines", headers={**env["h"], "Content-Type": "application/octet-stream"},
                       content=raw)
    result = resp.json()
    assert result["changed"] == 2 and [e["row"] for e in result["errors"]] == [4], result
    assert "не сходится" in result["errors"][0]["error"]
    a, b = line_of(result["run"], env["a"]), line_of(result["run"], env["b"])
    assert (a["accrued"], a["salary_amount"], a["bonus"], a["ndfl"], a["net"], a["manual_totals"]) == (
        "110000.00", "100000.00", "10000.00", "14300.00", "95000.00", True)
    assert a["insurance"] == "22595.93"  # МСП 2026: 30% с 1,5 МРОТ (12 191,85) + 15% сверху (10 404,08)
    assert (b["accrued"], b["ndfl"], b["insurance"], b["net"]) == ("30000.00", "3900.00", "9000.00", "26100.00")

    # норма дней меняется — готовые суммы не пересчитываются
    run = client.patch(f"/api/payroll-runs/{run['id']}", headers=env["h"], json={"norm_days": 18}).json()
    assert line_of(run, env["a"])["accrued"] == "110000.00"
    # проведение берёт загруженные суммы
    run = client.post(f"/api/payroll-runs/{run['id']}/post", headers=env["h"]).json()
    op = client.get(f"/api/operations/{run['operation_id']}", headers=env["h"]).json()
    assert sum(Decimal(p["amount"]) for p in op["postings"] if p["credit_account_id"] == env["acc"]["68.01"]) == Decimal("18200")
    client.post(f"/api/payroll-runs/{run['id']}/unpost", headers=env["h"])

    # правка вручную — строка снова считается программой
    run = client.patch(f"/api/payroll-lines/{a['id']}", headers=env["h"], json={"bonus": "0"}).json()
    a = line_of(run, env["a"])
    assert (a["manual_totals"], a["accrued"]) == (False, "100000.00")

    # шаблон для готовых сумм
    import io

    from openpyxl import load_workbook

    resp = client.get(f"/api/payroll-runs/{run['id']}/lines.xlsx", headers=env["h"], params={"ready": "true"})
    assert [c.value for c in load_workbook(io.BytesIO(resp.content)).active[1]][-4:] == [
        "Начислено", "НДФЛ", "Взносы", "К выплате"]


# --- Ведомости на выплату ---

from tests.test_bank_import import BANK, OUR, statement  # noqa: E402
from tests.test_calendar import expected as calendar_expected  # noqa: E402


@pytest.fixture
def bank_env(client, env, db):
    from app.models import Bank, Currency

    bank = Bank(bic="044525225", name="Сбер")
    db.add(bank)
    db.commit()
    rub = db.query(Currency).filter_by(code="RUB").one()
    acc = client.post(f"{env['o']}/bank-accounts", headers=env["h"], json={
        "legal_entity_id": env["le"].id, "bank_id": bank.id, "currency_id": rub.id, "number": OUR, "name": "Основной"})
    assert acc.status_code == 201, acc.text
    env["bank_acc"] = acc.json()
    return env


def sheet(client, env, kind, month=2, **extra):
    resp = client.post(f"{env['o']}/payout-sheets", headers=env["h"], json={
        "legal_entity_id": env["le"].id, "bank_account_id": env["bank_acc"]["id"], "kind": kind, "year": 2026,
        "month": month, "pay_date": f"2026-{month:02d}-15" if kind == "advance" else f"2026-{month + 1:02d}-05", **extra})
    assert resp.status_code == 201, resp.text
    return resp.json()


def us():
    return (OUR, 'ООО "А"', "7701234567")


def test_advance_confirmed_by_statement(client, bank_env):
    env = bank_env
    adv = sheet(client, env, "advance", advance_percent="40")
    amounts = {line["employee_name"]: line["amount"] for line in adv["lines"]}
    assert amounts == {"Андреев А.": "40000.00", "Борисов Б.": "24000.00"}
    b = next(line for line in adv["lines"] if line["employee_name"] == "Борисов Б.")
    adv = client.delete(f"/api/payout-lines/{b['id']}", headers=env["h"]).json()  # принят после аванса
    assert adv["total"] == "40000.00" and adv["number"] == 1

    reg = client.get(f"/api/payout-sheets/{adv['id']}/registry.xlsx", headers=env["h"])
    assert reg.status_code == 200
    sent = client.post(f"/api/payout-sheets/{adv['id']}/send", headers=env["h"]).json()
    assert sent["status"] == "sent"
    assert client.patch(f"/api/payout-lines/{adv['lines'][0]['id']}", headers=env["h"],
                        json={"amount": "1"}).status_code == 422  # отправленную не меняем

    raw = statement(OUR, [("501", "16.02.2026", "40000", us(), BANK, "Перечисление аванса по реестру № 1 за февраль")],
                    opening="100000", date_from="01.02.2026", date_to="28.02.2026")
    result = client.post(f"{env['o']}/bank-statements", headers={**env["h"], "Content-Type": "application/octet-stream"},
                         content=raw).json()
    assert result["confirmed_sheets"] == [1] and len(result["created"]) == 1
    paid = client.get(f"/api/payout-sheets/{adv['id']}", headers=env["h"]).json()
    assert (paid["status"], paid["paid_on"], paid["confirmed_by_statement"]) == ("paid", "2026-02-16", True)
    op = client.get(f"/api/operations/{paid['operation_id']}", headers=env["h"]).json()
    assert op["status"] == "posted" and op["postings"][0]["debit"] == {"employee": env["a"]["id"]}
    # повторная загрузка той же выписки — дубль
    again = client.post(f"{env['o']}/bank-statements", headers={**env["h"], "Content-Type": "application/octet-stream"},
                        content=raw).json()
    assert again["skipped_duplicates"] == 1 and again["confirmed_sheets"] == []


def test_salary_sheet_minus_advance_and_manual_pay(client, bank_env):
    env = bank_env
    adv = sheet(client, env, "advance")
    client.post(f"/api/payout-sheets/{adv['id']}/send", headers=env["h"])
    resp = client.post(f"{env['o']}/payout-sheets", headers=env["h"], json={
        "legal_entity_id": env["le"].id, "bank_account_id": env["bank_acc"]["id"], "kind": "salary",
        "year": 2026, "month": 2, "pay_date": "2026-03-05"})
    assert resp.status_code == 422 and "начисление" in resp.json()["detail"]

    run = create(client, env)
    assert client.post(f"/api/payroll-runs/{run['id']}/post", headers=env["h"]).status_code == 200
    run = client.get(f"/api/payroll-runs/{run['id']}", headers=env["h"]).json()
    sal = sheet(client, env, "salary")
    a_net = Decimal(line_of(run, env["a"])["net"])
    amounts = {line["employee_id"]: Decimal(line["amount"]) for line in sal["lines"]}
    assert amounts[env["a"]["id"]] == a_net - Decimal("50000.00")  # минус аванс 50%
    client.post(f"/api/payout-sheets/{sal['id']}/send", headers=env["h"])
    paid = client.post(f"/api/payout-sheets/{sal['id']}/pay", headers=env["h"], json={"paid_on": "2026-03-05"}).json()
    assert paid["status"] == "paid" and paid["confirmed_by_statement"] is False
    unpaid = client.post(f"/api/payout-sheets/{sal['id']}/unpay", headers=env["h"]).json()
    assert unpaid["status"] == "sent" and unpaid["operation_id"] is None
    assert client.post(f"/api/payout-sheets/{sal['id']}/unpay", headers=env["h"]).status_code == 422


def test_sheet_access(client, bank_env, db):
    env = bank_env
    adv = sheet(client, env, "advance")
    assert client.get(f"/api/payout-sheets/{adv['id']}", headers=env["ceo"]).status_code == 403
    assert client.get(f"/api/payout-sheets/{adv['id']}", headers=env["other"]).status_code == 404
    assert client.get(f"{env['o']}/payout-sheets", headers=env["ceo"]).status_code == 403
    # счёт чужого юрлица не подходит
    resp = client.post(f"{env['o']}/payout-sheets", headers=env["h"], json={
        "legal_entity_id": env["le2"].id, "bank_account_id": env["bank_acc"]["id"], "kind": "advance",
        "year": 2026, "month": 2, "pay_date": "2026-02-15"})
    assert resp.status_code == 422
    assert client.delete(f"/api/payout-sheets/{adv['id']}", headers=env["h"]).status_code == 204


def test_vacation_by_average_earnings(client, env):
    for month in (1, 2, 3):  # три полных месяца по 100 000
        run = create(client, env, month)
        assert client.post(f"/api/payroll-runs/{run['id']}/post", headers=env["h"]).status_code == 200
    run = create(client, env, 4)
    a = line_of(run, env["a"])
    resp = client.post(f"/api/payroll-lines/{a['id']}/vacation", headers=env["h"], json={"days": 14})
    assert resp.status_code == 200, resp.text
    a = line_of(resp.json(), env["a"])
    # 300 000 / (3 × 29,3) = 3 412,97 в день × 14
    assert a["vacation"] == "47781.58"
    assert "средний 3 412,97" in a["vacation_note"] and "87,90 дн." in a["vacation_note"]
    assert a["accrued"] == str(Decimal(a["salary_amount"]) + Decimal("47781.58"))

    # без истории — оклад / 29,3
    b = line_of(run, env["b"])
    b = line_of(client.post(f"/api/payroll-lines/{b['id']}/vacation", headers=env["h"], json={"days": 7}).json(), env["b"])
    assert b["vacation"] == str((Decimal("60000") / Decimal("29.3")).quantize(Decimal("0.01")) * 7)

    # сумма, поправленная вручную, сбрасывает пояснение
    run = client.patch(f"/api/payroll-lines/{a['id']}", headers=env["h"], json={"vacation": "40000"}).json()
    assert line_of(run, env["a"])["vacation_note"] == ""


def test_sick_by_average_earnings(client, env):
    run = create(client, env, 4)
    a = line_of(run, env["a"])
    insurance_before = a["insurance"]
    resp = client.post(f"/api/payroll-lines/{a['id']}/sick", headers=env["h"], json={
        "days": 10, "experience": "5to8", "earnings": {"2024": "1200000", "2025": "3000000"}})
    assert resp.status_code == 200, resp.text
    a = line_of(resp.json(), env["a"])
    # (1 200 000 + 2 759 000 предельная база 2025) / 730 = 5 423,29 × 80% = 4 338,63; работодатель — 3 дня
    assert a["sick"] == "13015.89"
    assert "СФР 7 дн. = 30 370,41" in a["sick_note"]
    assert a["insurance"] == insurance_before  # пособие взносами не облагается

    # меньше 6 месяцев стажа — из МРОТ
    a = line_of(client.post(f"/api/payroll-lines/{a['id']}/sick", headers=env["h"],
                            json={"days": 2, "experience": "lt6m"}).json(), env["a"])
    assert a["sick"] == str((Decimal("27093") * 24 / 730).quantize(Decimal("0.01")) * 2)

    # проведённое начисление не пересчитывается
    client.post(f"/api/payroll-runs/{run['id']}/post", headers=env["h"])
    resp = client.post(f"/api/payroll-lines/{a['id']}/sick", headers=env["h"], json={"days": 2, "experience": "ge8"})
    assert resp.status_code == 422
    resp = client.post(f"/api/payroll-lines/{a['id']}/vacation", headers=env["ceo"], json={"days": 2})
    assert resp.status_code == 403


def test_payout_sheets(client, bank_env):
    env = bank_env
    adv = sheet(client, env, "advance", advance_percent="40")
    feb = {"date_from": "2026-02-01", "date_to": "2026-02-28"}
    data = calendar_expected(client, env, **feb)
    assert data["brief"] == [("payout", "2026-02-15", "out", "64000.00")]
    assert data["rows"][0]["title"] == "Ведомость № 1 (аванс 02.2026)"
    client.post(f"/api/payout-sheets/{adv['id']}/send", headers=env["h"])
    assert len(calendar_expected(client, env, **feb)["rows"]) == 1
    client.post(f"/api/payout-sheets/{adv['id']}/pay", headers=env["h"], json={"paid_on": "2026-02-15"})
    assert calendar_expected(client, env, **feb)["rows"] == []
