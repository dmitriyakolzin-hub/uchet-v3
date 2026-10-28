"""Сетка ввода плана: ячейка = плановая операция, копирование версии, права (docs/TZ.md, раздел 13)."""

import pytest

from tests.test_users import admin  # noqa: F401


@pytest.fixture
def env(client, headers_for, make_user, make_org, db, admin):  # noqa: F811
    from app.models import LegalEntity

    fd = make_user(login="fd")
    ceo = make_user(login="ceo")
    head = make_user(login="head")
    org = make_org("Группа А", members=[(fd, "finance_director"), (ceo, "ceo")])
    le = LegalEntity(organization_id=org.id, full_name="ООО А", short_name="А", inn="7701234567")
    db.add(le)
    db.commit()
    h = headers_for(fd)
    o = f"/api/organizations/{org.id}"

    def post(path, body):
        resp = client.post(f"{o}/{path}", headers=h, json=body)
        assert resp.status_code == 201, resp.text
        return resp.json()

    company = post("cfos", {"name": "Компания", "legal_entity_id": le.id})
    trade = post("cfos", {"name": "Торговля", "parent_id": company["id"]})
    shop = post("cfos", {"name": "Магазин 1", "parent_id": trade["id"]})
    logistics = post("cfos", {"name": "Логистика", "parent_id": company["id"]})
    client.put(f"/api/users/{head.id}/memberships", headers=headers_for(admin), json={
        "memberships": [{"organization_id": org.id, "role": "cfo_head", "cfo_id": trade["id"]}]})
    budget = post("budgets", {"kind": "year", "year": 2027})
    art = {a["name"]: a["id"] for a in client.get(f"/api/organizations/{org.id}/articles", headers=h).json()}
    acc = {a["id"]: a["code"] for a in client.get("/api/accounts", headers=h).json()}
    return {"org": org, "le": le, "o": o, "h": h, "ceo": headers_for(ceo), "head": headers_for(head),
            "company": company, "trade": trade, "shop": shop, "logistics": logistics, "budget": budget,
            "version": budget["versions"][0], "art": art, "acc": acc}


def put(client, env, cells, kind="pl", cfo=None, h=None, version=None):
    cfo_id = env[cfo]["id"] if cfo else None
    return client.put(f"/api/budget-versions/{(version or env['version'])['id']}/cells", headers=h or env["h"],
                      json={"kind": kind, "cfo_id": cfo_id, "cells": cells})


def grid(client, env, kind="pl", cfo="shop", h=None, version=None):
    resp = client.get(f"/api/budget-versions/{(version or env['version'])['id']}/grid", headers=h or env["h"],
                      params={"kind": kind, **({"cfo_id": env[cfo]["id"]} if cfo else {})})
    return resp


def test_cells_become_plan_operations(client, env, db):
    from app.models import Operation

    rent, revenue = env["art"]["Аренда"], env["art"]["Выручка от продаж"]
    resp = put(client, env, [{"article_id": rent, "month": 1, "amount": "1000"},
                             {"article_id": rent, "month": 2, "amount": "1000.5"},
                             {"article_id": revenue, "month": 1, "amount": "50000"}], cfo="shop")
    assert resp.status_code == 200, resp.text
    assert resp.json() == {"changed": 3}

    data = grid(client, env).json()
    assert data["editable"] is True and data["budget_title"] == "Бюджет 2027"
    assert [m["key"] for m in data["months"]][:2] == ["2027-01", "2027-02"] and len(data["months"]) == 12
    rows = {r["article_id"]: r for r in data["rows"]}
    assert rows[rent]["values"] == {"2027-01": "1000.00", "2027-02": "1000.50"}
    assert rows[rent]["group"] == "Коммерческие расходы"
    assert data["rows"][0]["group"] == "Выручка"  # порядок — как в ОПиУ

    op = db.query(Operation).filter_by(article_id=rent, plan_cfo_id=env["shop"]["id"]).order_by(Operation.op_date).first()
    assert (op.status, op.number, op.legal_entity_id, str(op.op_date)) == ("plan", 0, env["le"].id, "2027-01-01")
    p = op.postings[0]
    assert (env["acc"][p.debit_account_id], env["acc"][p.credit_account_id]) == ("44", "60.01")
    assert (p.dt_cfo_id, p.dt_pl_article_id) == (env["shop"]["id"], rent)

    # изменение и очистка ячейки
    put(client, env, [{"article_id": rent, "month": 1, "amount": "1200"}, {"article_id": rent, "month": 2}], cfo="shop")
    rows = {r["article_id"]: r for r in grid(client, env).json()["rows"]}
    assert rows[rent]["values"] == {"2027-01": "1200.00"}
    # у другого ЦФО — своя сетка
    assert {r["article_id"]: r for r in grid(client, env, cfo="logistics").json()["rows"]}[rent]["values"] == {}


def test_plan_invisible_in_journal_and_reports(client, env):
    put(client, env, [{"article_id": env["art"]["Аренда"], "month": 1, "amount": "1000"}], cfo="shop")
    items = client.get(f"{env['o']}/operations", headers=env["h"]).json()["items"]
    assert items == []
    osv = client.get(f"{env['o']}/reports/osv", headers=env["h"],
                     params={"date_from": "2027-01-01", "date_to": "2027-12-31", "drafts": True}).json()
    assert osv["rows"] == []


def test_cell_validation(client, env):
    rent = env["art"]["Аренда"]
    assert put(client, env, [{"article_id": rent, "month": 1, "amount": "-5"}], cfo="shop").status_code == 422
    assert put(client, env, [{"article_id": rent, "month": 1, "amount": "5"}]).status_code == 422  # ОПиУ без ЦФО
    assert put(client, env, [{"article_id": env["art"]["Оплата поставщикам"], "month": 1, "amount": "5"}],
               cfo="shop").status_code == 422  # статья ДДС в плане ОПиУ
    assert put(client, env, [{"article_id": rent, "month": 13, "amount": "5"}], cfo="shop").status_code == 422
    q = client.post(f"{env['o']}/budgets", headers=env["h"], json={"kind": "quarter", "year": 2027, "period": 2}).json()
    resp = put(client, env, [{"article_id": rent, "month": 1, "amount": "5"}], cfo="shop", version=q["versions"][0])
    assert resp.status_code == 422 and "не входит в период" in resp.json()["detail"]
    assert put(client, env, [{"article_id": rent, "month": 4, "amount": "5"}], cfo="shop",
               version=q["versions"][0]).status_code == 200


def test_cash_plan_without_cfo(client, env):
    pay = env["art"]["Оплата поставщикам"]
    assert put(client, env, [{"article_id": pay, "month": 3, "amount": "700"}], kind="cf").status_code == 200
    data = grid(client, env, kind="cf", cfo=None).json()
    rows = {r["article_id"]: r for r in data["rows"]}
    assert rows[pay]["values"] == {"2027-03": "700.00"} and rows[pay]["group"] == "Операционная: выплаты"
    assert env["art"]["Перевод между своими счетами"] not in rows


def test_approved_version_is_frozen_and_copied(client, env):
    rent = env["art"]["Аренда"]
    put(client, env, [{"article_id": rent, "month": 1, "amount": "1000"}], cfo="shop")
    put(client, env, [{"article_id": env["art"]["Оплата поставщикам"], "month": 1, "amount": "900"}], kind="cf")
    assert client.post(f"/api/budget-versions/{env['version']['id']}/approve", headers=env["h"]).status_code == 200
    resp = put(client, env, [{"article_id": rent, "month": 1, "amount": "5"}], cfo="shop")
    assert resp.status_code == 422 and "черновик" in resp.json()["detail"]
    assert grid(client, env).json()["editable"] is False

    v2 = client.post(f"/api/budgets/{env['budget']['id']}/versions", headers=env["h"],
                     json={"copy_from_id": env["version"]["id"]}).json()
    rows = {r["article_id"]: r for r in grid(client, env, version=v2).json()["rows"]}
    assert rows[rent]["values"] == {"2027-01": "1000.00"}
    cf = {r["article_id"]: r for r in grid(client, env, kind="cf", cfo=None, version=v2).json()["rows"]}
    assert cf[env["art"]["Оплата поставщикам"]]["values"] == {"2027-01": "900.00"}
    # правка новой версии не меняет утверждённую
    put(client, env, [{"article_id": rent, "month": 1, "amount": "1500"}], cfo="shop", version=v2)
    assert {r["article_id"]: r for r in grid(client, env).json()["rows"]}[rent]["values"] == {"2027-01": "1000.00"}
    # удаление черновика удаляет его план
    assert client.delete(f"/api/budget-versions/{v2['id']}", headers=env["h"]).status_code == 204


def test_plan_permissions(client, env):
    rent = env["art"]["Аренда"]
    assert put(client, env, [{"article_id": rent, "month": 1, "amount": "5"}], cfo="shop", h=env["ceo"]).status_code == 403
    assert grid(client, env, h=env["ceo"]).json()["editable"] is False
    # руководитель ЦФО — только своя ветка
    assert put(client, env, [{"article_id": rent, "month": 1, "amount": "5"}], cfo="shop", h=env["head"]).status_code == 200
    assert put(client, env, [{"article_id": rent, "month": 1, "amount": "5"}], cfo="logistics",
               h=env["head"]).status_code == 404
    assert grid(client, env, cfo="logistics", h=env["head"]).status_code == 404
    assert grid(client, env, kind="cf", cfo=None, h=env["head"]).status_code == 404
    assert grid(client, env, h=env["head"]).json()["editable"] is True


def test_plan_excel_roundtrip(client, env):
    import io

    from openpyxl import load_workbook

    rent = env["art"]["Аренда"]
    put(client, env, [{"article_id": rent, "month": 1, "amount": "1000"}], cfo="shop")
    resp = client.get(f"/api/budget-versions/{env['version']['id']}/plan.xlsx", headers=env["h"], params={"kind": "pl"})
    assert resp.status_code == 200
    wb = load_workbook(io.BytesIO(resp.content))
    ws = wb.active
    assert [c.value for c in ws[1]][:5] == ["ЦФО", "Статья", "ЦФУ", "Январь", "Февраль"]
    rows = {(r[0], r[1]): r for r in ws.iter_rows(min_row=2, values_only=True)}
    assert rows[("Компания / Торговля / Магазин 1", "Аренда")][3] == 1000
    assert ("Компания / Логистика", "Аренда") in rows

    # правим файл: январь очищаем, февраль и март ставим; строка с ошибкой не грузится
    for row in ws.iter_rows(min_row=2):
        if (row[0].value, row[1].value) == ("Компания / Торговля / Магазин 1", "Аренда"):
            row[3].value, row[4].value, row[5].value = None, 2000, "1 500,50"
    ws.append(["Нет такого", "Аренда", None, 5])
    ws.append(["Логистика", "Аренда", None, 0, 300])  # ЦФО можно указать и просто названием
    out = io.BytesIO()
    wb.save(out)
    resp = client.post(f"/api/budget-versions/{env['version']['id']}/plan", headers={**env["h"], "Content-Type":
                       "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"},
                       params={"kind": "pl"}, content=out.getvalue())
    assert resp.status_code == 200, resp.text
    result = resp.json()
    assert result["changed"] == 4 and len(result["errors"]) == 1 and "Нет такого" in result["errors"][0]["error"]
    rows = {r["article_id"]: r for r in grid(client, env).json()["rows"]}
    assert rows[rent]["values"] == {"2027-02": "2000.00", "2027-03": "1500.50"}
    rows = {r["article_id"]: r for r in grid(client, env, cfo="logistics").json()["rows"]}
    assert rows[rent]["values"] == {"2027-02": "300.00"}

    # руководитель ЦФО скачивает только свою ветку и не грузит чужую
    head = client.get(f"/api/budget-versions/{env['version']['id']}/plan.xlsx", headers=env["head"])
    labels = {r[0] for r in load_workbook(io.BytesIO(head.content)).active.iter_rows(min_row=2, values_only=True)}
    assert labels == {"Компания / Торговля", "Компания / Торговля / Магазин 1"}
    resp = client.post(f"/api/budget-versions/{env['version']['id']}/plan", headers={**env["head"], "Content-Type":
                       "application/octet-stream"}, content=out.getvalue())
    assert all("ветки" in e["error"] for e in resp.json()["errors"] if "Нет такого" not in e["error"])


def branch(client, env, cfo, h=None):
    return client.get(f"/api/budget-versions/{env['version']['id']}/grid", headers=h or env["h"],
                      params={"kind": "pl", "cfo_id": env[cfo]["id"], "mode": "branch"})


def test_branch_rollup(client, env):
    rent = env["art"]["Аренда"]
    put(client, env, [{"article_id": rent, "month": 1, "amount": "1000"}], cfo="shop")
    put(client, env, [{"article_id": rent, "month": 1, "amount": "300"}], cfo="trade")  # свой план направления
    put(client, env, [{"article_id": rent, "month": 2, "amount": "50"}], cfo="logistics")

    data = branch(client, env, "company").json()
    assert (data["mode"], data["has_children"], data["editable"]) == ("branch", True, False)
    row = next(r for r in data["rows"] if r["article_id"] == rent)
    assert row["values"] == {"2027-01": "1300.00", "2027-02": "50.00"}
    assert [(p["name"], p["values"]) for p in row["parts"]] == [
        ("Логистика", {"2027-02": "50.00"}), ("Торговля", {"2027-01": "1300.00"})]

    row = next(r for r in branch(client, env, "trade").json()["rows"] if r["article_id"] == rent)
    assert [(p["name"], p["values"]) for p in row["parts"]] == [
        ("Торговля — свой план", {"2027-01": "300.00"}), ("Магазин 1", {"2027-01": "1000.00"})]

    # свой план и лист без подчинённых — как раньше
    own = grid(client, env, cfo="trade").json()
    assert own["mode"] == "own" and own["editable"] and own["has_children"]
    assert next(r for r in own["rows"] if r["article_id"] == rent)["values"] == {"2027-01": "300.00"}
    assert grid(client, env, cfo="shop").json()["has_children"] is False

    # руководитель ЦФО видит свод своей ветки, но не выше
    assert branch(client, env, "trade", h=env["head"]).status_code == 200
    assert branch(client, env, "company", h=env["head"]).status_code == 404

    resp = client.get(f"/api/budget-versions/{env['version']['id']}/branch.xlsx", headers=env["h"],
                      params={"cfo_id": env["company"]["id"]})
    assert resp.status_code == 200 and resp.content[:2] == b"PK"
