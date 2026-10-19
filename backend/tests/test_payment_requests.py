"""Заявки на оплату: согласование, права, контроль бюджета ДДС, подтверждение выпиской (docs/TZ.md, раздел 18)."""

import io

import pytest
from openpyxl import load_workbook

from tests.test_bank_import import BANK, OUR, statement
from tests.test_closing import env as closing_env  # noqa: F401

SUPPLIER = (BANK[0], "ООО Арендодатель", "7702222222")


@pytest.fixture
def env(client, closing_env, db):  # noqa: F811
    from app.models import Bank, Currency, Membership, User

    env = closing_env
    bank = Bank(bic="044525225", name="Сбер")
    db.add(bank)
    db.commit()
    rub = db.query(Currency).filter_by(code="RUB").one()
    acc = client.post(f"{env['o']}/bank-accounts", headers=env["h"], json={
        "legal_entity_id": env["le"].id, "bank_id": bank.id, "currency_id": rub.id, "number": OUR, "name": "Основной"})
    assert acc.status_code == 201, acc.text
    env["bank_acc"] = acc.json()
    env["shop"] = env["post"]("cfos", {"name": "Магазин"})
    head = db.query(User).filter_by(login="head").one()
    membership = db.query(Membership).filter_by(user_id=head.id).one()
    membership.cfo_id = env["office"]["id"]
    db.commit()
    return env


def create(client, env, h=None, **fields):
    body = {"pay_date": "2026-03-10", "legal_entity_id": env["le"].id, "counterparty_id": env["cp"]["id"],
            "contract_id": env["contract"]["id"], "article_id": env["art"]["Оплата поставщикам"],
            "cfo_id": env["office"]["id"], "amount": "30000", "description": "Аренда за март", **fields}
    return client.post(f"{env['o']}/payment-requests", headers=h or env["h"], json=body)


def act(client, env, req, action, h=None, **body):
    return client.post(f"/api/payment-requests/{req['id']}/{action}", headers=h or env["h"], json=body or None)


def test_workflow_and_access(client, env):
    # руководитель ЦФО подаёт заявку своей ветки, счёт выбирает финансовый директор
    assert create(client, env, h=env["head"], cfo_id=env["shop"]["id"]).status_code == 422
    assert create(client, env, h=env["head"], bank_account_id=env["bank_acc"]["id"]).status_code == 422
    resp = create(client, env, h=env["head"])
    assert resp.status_code == 201, resp.text
    req = resp.json()
    assert (req["number"], req["status"], req["source_name"]) == (1, "draft", "")
    assert create(client, env, h=env["ceo"]).status_code == 403
    assert create(client, env, h=env["other"]).status_code == 404

    req = act(client, env, req, "submit", h=env["head"]).json()
    assert req["status"] == "pending"
    assert act(client, env, req, "approve", h=env["head"]).status_code == 403
    assert act(client, env, req, "approve", h=env["ceo"]).status_code == 403
    assert client.put(f"/api/payment-requests/{req['id']}", headers=env["head"], json={
        "pay_date": "2026-03-11", "legal_entity_id": env["le"].id, "amount": "1",
        "cfo_id": env["office"]["id"]}).status_code == 422  # на согласовании не правится

    req = act(client, env, req, "reject", reason="Нет счёта").json()
    assert (req["status"], req["reject_reason"]) == ("rejected", "Нет счёта")
    req = act(client, env, req, "submit", h=env["head"]).json()
    req = act(client, env, req, "approve").json()
    assert req["status"] == "approved" and req["approved_by_name"]
    assert client.delete(f"/api/payment-requests/{req['id']}", headers=env["h"]).status_code == 422

    # перенос и выбор счёта без повторного согласования
    req = act(client, env, req, "move", pay_date="2026-03-12", bank_account_id=env["bank_acc"]["id"]).json()
    assert (req["status"], req["pay_date"], req["source_name"]) == ("approved", "2026-03-12", "Основной ··0001")
    head_view = client.get(f"/api/payment-requests/{req['id']}", headers=env["head"]).json()
    assert head_view["source_name"] == "" and head_view["bank_account_id"] is None

    # заявки чужой ветки руководитель ЦФО не видит
    other = create(client, env, cfo_id=env["shop"]["id"]).json()
    listed = client.get(f"{env['o']}/payment-requests", headers=env["head"]).json()
    assert [r["number"] for r in listed["rows"]] == [1]
    assert client.get(f"/api/payment-requests/{other['id']}", headers=env["head"]).status_code == 404
    assert client.get(f"/api/payment-requests/{other['id']}", headers=env["other"]).status_code == 404
    assert client.get(f"{env['o']}/payment-requests", headers=env["ceo"]).json()["total_out"] == "60000.00"
    assert client.delete(f"/api/payment-requests/{other['id']}", headers=env["h"]).status_code == 204

    # ожидаемое поступление согласования не требует
    income = create(client, env, direction="in", article_id=env["art"]["Поступление от покупателей"]).json()
    assert act(client, env, income, "submit").json()["status"] == "approved"
    assert create(client, env, direction="in").status_code == 422  # статья выплат у поступления

    # ручная отметка оплаты и её снятие
    assert act(client, env, req, "pay", paid_on="2026-03-12", operation_number=999).status_code == 422
    paid = act(client, env, req, "pay", paid_on="2026-03-12").json()
    assert (paid["status"], paid["paid_on"], paid["confirmed_by_statement"]) == ("paid", "2026-03-12", False)
    assert act(client, env, paid, "approve").status_code == 422
    assert act(client, env, paid, "unpay").json()["status"] == "approved"

    xlsx = client.get(f"{env['o']}/payment-requests.xlsx", headers=env["h"], params={"status": "approved"})
    assert xlsx.status_code == 200
    ws = load_workbook(io.BytesIO(xlsx.content)).active
    numbers = [row[0] for row in ws.iter_rows(min_row=6, values_only=True)]
    assert numbers == [2, 1, "Итого"]  # по дате оплаты: поступление 10.03, выплата 12.03


def test_budget_control(client, env):
    o, h = env["o"], env["h"]
    pay = env["art"]["Оплата поставщикам"]
    unbudgeted = create(client, env).json()
    assert unbudgeted["budget"]["plan"] is None and unbudgeted["budget"]["over"] == "0.00"

    budget = client.post(f"{o}/budgets", headers=h, json={"kind": "year", "year": 2026}).json()
    version = budget["versions"][0]
    resp = client.put(f"/api/budget-versions/{version['id']}/cells", headers=h, json={
        "kind": "cf", "cells": [{"article_id": pay, "month": 3, "amount": "100000"}]})
    assert resp.status_code == 200, resp.text
    assert client.post(f"/api/budget-versions/{version['id']}/approve", headers=h).status_code == 200

    # факт: оплачено 50 000 в марте
    resp = client.post(f"{o}/operations?post=true", headers=h, json={
        "legal_entity_id": env["le"].id, "op_date": "2026-03-02", "postings": [{
            "debit_account_id": env["acc"]["60.01"], "credit_account_id": env["acc"]["51"], "amount": "50000",
            "debit": {"counterparty": env["cp"]["id"], "contract": env["contract"]["id"]},
            "credit": {"bank_account": env["bank_acc"]["id"], "cash_flow_item": pay}}]})
    assert resp.status_code == 201, resp.text

    first = act(client, env, unbudgeted, "submit").json()
    assert first["budget"] == {"month": "2026-03", "plan": "100000.00", "fact": "50000.00", "requested": "0.00",
                               "remaining": "50000.00", "over": "0.00"}
    second = create(client, env, amount="25000").json()
    assert (second["budget"]["requested"], second["budget"]["remaining"], second["budget"]["over"]) == \
        ("30000.00", "20000.00", "5000.00")  # сверх бюджета на 5 000 — предупреждение, не запрет
    assert act(client, env, second, "submit").json()["status"] == "pending"


def test_statement_confirms_request(client, env):
    req = create(client, env, pay_date="2026-03-05").json()
    act(client, env, req, "submit")
    act(client, env, req, "approve")
    other = create(client, env, pay_date="2026-03-25").json()  # та же сумма, дальше по дате
    act(client, env, other, "submit")
    act(client, env, other, "approve")

    us = (OUR, 'ООО "А"', "7701234567")
    raw = statement(OUR, [("71", "06.03.2026", "30000", us, SUPPLIER, "Оплата по счёту 15 за аренду"),
                          ("72", "06.03.2026", "777", us, SUPPLIER, "Прочее")],
                    opening="100000", date_from="01.03.2026", date_to="31.03.2026")
    result = client.post(f"{env['o']}/bank-statements", headers={**env["h"], "Content-Type": "application/octet-stream"},
                         content=raw).json()
    assert result["paid_requests"] == [1] and len(result["created"]) == 2

    paid = client.get(f"/api/payment-requests/{req['id']}", headers=env["h"]).json()
    assert (paid["status"], paid["paid_on"], paid["confirmed_by_statement"]) == ("paid", "2026-03-06", True)
    op = client.get(f"/api/operations/{paid['operation_id']}", headers=env["h"]).json()
    assert op["article_id"] == env["art"]["Оплата поставщикам"] and "заявка № 1" in op["description"]
    assert op["postings"][0]["debit"] == {"counterparty": env["cp"]["id"], "contract": env["contract"]["id"]}
    assert client.get(f"/api/payment-requests/{other['id']}", headers=env["h"]).json()["status"] == "approved"
