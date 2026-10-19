"""Загрузка операций и контрагентов из Excel (docs/TZ.md, раздел 11, часть 1.7б)."""

import datetime as dt
import io

import pytest
from openpyxl import Workbook, load_workbook

OUR_INN = "7701234567"
BANK_NUMBER = "40702810900000000001"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

OP_HEADER = ["Дата*", "Операция", "Юрлицо (ИНН)", "Статья", "Счёт Дт*", "Счёт Кт*", "Сумма*", "Валюта",
             "Количество", "Контрагент (ИНН или наименование)", "Договор №", "Расчётный счёт (номер)",
             "Сотрудник (таб. № или ФИО)", "Номенклатура (артикул или наименование)", "Склад",
             "Основное средство (инв. № или наименование)", "ЦФО", "Описание"]


def xlsx(header, rows) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append(header)
    for row in rows:
        ws.append(row)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def op_row(**cols):
    """Строка шаблона операций по названиям колонок без звёздочек."""
    names = [h.rstrip("*") for h in OP_HEADER]
    return [cols.get(n) for n in names]


@pytest.fixture
def env(client, headers_for, make_user, make_org, db):
    from app.models import Bank, Cfo, Currency, LegalEntity

    fd = make_user(login="fd")
    ceo = make_user(login="ceo")
    other_fd = make_user(login="other")
    org = make_org("Группа А", members=[(fd, "finance_director"), (ceo, "ceo")])
    make_org("Клиент Б", kind="client", members=[(other_fd, "finance_director")])
    le = LegalEntity(organization_id=org.id, full_name="ООО Наша", short_name="Наша", inn=OUR_INN)
    bank = Bank(bic="044525225", name="Банк")
    db.add_all([le, bank])
    db.commit()
    h = headers_for(fd)
    o = f"/api/organizations/{org.id}"
    rub = db.query(Currency).filter_by(code="RUB").one()
    ba = client.post(f"{o}/bank-accounts", headers=h, json={
        "legal_entity_id": le.id, "bank_id": bank.id, "currency_id": rub.id, "number": BANK_NUMBER,
        "name": "Основной"}).json()
    cp = client.post(f"{o}/counterparties", headers=h, json={
        "name": "ООО Поставщик", "kind": "legal", "inn": "7702222222", "is_supplier": True}).json()
    contract = client.post(f"{o}/contracts", headers=h, json={
        "counterparty_id": cp["id"], "legal_entity_id": le.id, "currency_id": rub.id, "kind": "supplier",
        "number": "П-7"}).json()
    company = client.post(f"{o}/cfos", headers=h, json={"name": "Компания", "legal_entity_id": le.id}).json()
    shop = client.post(f"{o}/cfos", headers=h, json={"name": "Офис", "parent_id": company["id"]}).json()
    assert db.query(Cfo).count() == 2
    accounts = {a["code"]: a["id"] for a in client.get("/api/accounts", headers=h).json()}
    return {"org": org, "le": le, "o": o, "h": h, "ba": ba, "cp": cp, "contract": contract, "shop": shop,
            "acc": accounts, "ceo": ceo, "other_fd": other_fd, "headers_for": headers_for}


def upload(client, env, path, raw, h=None):
    return client.post(f"{env['o']}/import/{path}", headers={**(h or env["h"]), "Content-Type": XLSX}, content=raw)


def operations(client, env):
    items = client.get(f"{env['o']}/operations", headers=env["h"], params={"limit": 1000}).json()["items"]
    return sorted(items, key=lambda o: o["number"])


PAYMENT = dict(Дата=dt.date(2026, 10, 7), Статья="Оплата поставщикам", **{
    "Счёт Дт": "60.01", "Счёт Кт": "51", "Сумма": 2400.5, "Контрагент (ИНН или наименование)": "7702222222",
    "Договор №": "П-7", "Расчётный счёт (номер)": BANK_NUMBER, "Описание": "Оплата по счёту 15"})


def test_operations_from_template_become_drafts(client, env):
    rent = dict(Дата="10.10.2026", Статья="Аренда", **{
        "Счёт Дт": 26, "Счёт Кт": 60.01, "Сумма": "15 000,00", "ЦФО": "Офис",
        "Контрагент (ИНН или наименование)": "ООО Поставщик", "Договор №": "П-7"})
    resp = upload(client, env, "operations", xlsx(OP_HEADER, [op_row(**PAYMENT), op_row(**rent)]))
    assert resp.status_code == 200, resp.text
    result = resp.json()
    assert result == {"rows": 2, "created": [1, 2], "incomplete": [], "skipped_duplicates": 0, "errors": []}

    pay, cost = operations(client, env)
    assert pay["status"] == "draft" and pay["op_date"] == "2026-10-07" and pay["description"] == "Оплата по счёту 15"
    p = pay["postings"][0]
    assert (p["debit_account_id"], p["credit_account_id"], p["amount"]) == (env["acc"]["60.01"], env["acc"]["51"], "2400.50")
    assert p["debit"] == {"counterparty": env["cp"]["id"], "contract": env["contract"]["id"]}
    assert p["credit"] == {"bank_account": env["ba"]["id"], "cash_flow_item": pay["article_id"]}
    c = cost["postings"][0]
    assert c["debit"] == {"cfo": env["shop"]["id"], "pl_item": cost["article_id"]}
    assert c["amount"] == "15000.00"
    # обе операции полные — проводятся
    for op in (pay, cost):
        assert client.post(f"/api/operations/{op['id']}/post", headers=env["h"]).status_code == 200


def test_rows_with_errors_are_reported_and_skipped(client, env):
    rows = [
        op_row(**PAYMENT),
        op_row(**{**PAYMENT, "Счёт Дт": "99.99"}),
        op_row(**{**PAYMENT, "Контрагент (ИНН или наименование)": "7703333333"}),
        op_row(**{**PAYMENT, "Сумма": "12.345"}),
        op_row(**{**PAYMENT, "Дата": "32.10.2026"}),
        op_row(**{**PAYMENT, "Склад": "Главный"}),
        op_row(**{**PAYMENT, "Договор №": "нет такого"}),
    ]
    result = upload(client, env, "operations", xlsx(OP_HEADER, rows)).json()
    assert result["created"] == [1]
    errors = {e["row"]: e["error"] for e in result["errors"]}
    assert set(errors) == {3, 4, 5, 6, 7, 8}
    assert "99.99" in errors[3]
    assert "7703333333" in errors[4]
    assert "2 знаков" in errors[5]
    assert "Дата" in errors[6]
    assert "Склад" in errors[7]  # склад не найден в справочнике
    assert "нет такого" in errors[8]


def test_unneeded_analytics_is_an_error(client, env, db):
    from app.models import Warehouse

    db.add(Warehouse(organization_id=env["org"].id, name="Главный"))
    db.commit()
    result = upload(client, env, "operations", xlsx(OP_HEADER, [op_row(**{**PAYMENT, "Склад": "Главный"})])).json()
    assert result["created"] == [] and "не ведётся" in result["errors"][0]["error"]


def test_incomplete_draft_is_created(client, env):
    # нет договора и расчётного счёта — черновик создаётся, но провести его пока нельзя
    row = {k: v for k, v in PAYMENT.items() if k not in ("Договор №", "Расчётный счёт (номер)")}
    result = upload(client, env, "operations", xlsx(OP_HEADER, [op_row(**row)])).json()
    assert result["created"] == [1] and result["incomplete"] == [1]


def test_grouped_rows_make_one_operation(client, env):
    sale = dict(Дата="15.10.2026", Операция="A", Статья="Выручка от продаж", **{
        "Контрагент (ИНН или наименование)": "7702222222", "Договор №": "П-7", "ЦФО": "Офис"})
    rows = [op_row(**sale, **{"Счёт Дт": "62.01", "Счёт Кт": "90.01", "Сумма": 1200}),
            op_row(**sale, **{"Счёт Дт": "90.03", "Счёт Кт": "68.02", "Сумма": 200})]
    result = upload(client, env, "operations", xlsx(OP_HEADER, rows)).json()
    assert result["created"] == [1], result
    (op,) = operations(client, env)
    assert [p["amount"] for p in op["postings"]] == ["1200.00", "200.00"]

    bad = [op_row(**sale, **{"Счёт Дт": "62.01", "Счёт Кт": "90.01", "Сумма": 1}),
           op_row(**{**sale, "Дата": "16.10.2026"}, **{"Счёт Дт": "90.03", "Счёт Кт": "68.02", "Сумма": 1})]
    result = upload(client, env, "operations", xlsx(OP_HEADER, bad)).json()
    assert result["created"] == [] and "строка 3" in result["errors"][0]["error"]


def test_reupload_creates_no_duplicates_but_keeps_repeated_rows(client, env):
    raw = xlsx(OP_HEADER, [op_row(**PAYMENT), op_row(**PAYMENT)])  # два одинаковых платежа — оба нужны
    assert upload(client, env, "operations", raw).json()["created"] == [1, 2]
    again = upload(client, env, "operations", raw).json()
    assert again["created"] == [] and again["skipped_duplicates"] == 2
    # дописали третий такой же платёж — загрузится только он
    raw3 = xlsx(OP_HEADER, [op_row(**PAYMENT)] * 3)
    third = upload(client, env, "operations", raw3).json()
    assert third["created"] == [3] and third["skipped_duplicates"] == 2


def test_legal_entity_required_when_several(client, env, db):
    from app.models import LegalEntity

    db.add(LegalEntity(organization_id=env["org"].id, full_name="ИП Второй", short_name="Второй", inn="770111111111"))
    db.commit()
    result = upload(client, env, "operations", xlsx(OP_HEADER, [op_row(**PAYMENT)])).json()
    assert "юрлицо" in result["errors"][0]["error"]
    result = upload(client, env, "operations", xlsx(OP_HEADER, [op_row(**PAYMENT, **{"Юрлицо (ИНН)": OUR_INN})])).json()
    assert result["created"] == [1]


def test_bad_files(client, env):
    resp = upload(client, env, "operations", b"not an excel file")
    assert resp.status_code == 422 and ".xlsx" in resp.json()["detail"]
    resp = upload(client, env, "operations", b"PK\x03\x04broken")
    assert resp.status_code == 422
    resp = upload(client, env, "operations", xlsx(["Что-то", "ещё"], [[1, 2]]))
    assert resp.status_code == 422 and "«Дата»" in resp.json()["detail"]
    assert upload(client, env, "operations", b"").status_code == 422


def test_operations_template(client, env):
    resp = client.get(f"{env['o']}/import/templates/operations.xlsx", headers=env["h"])
    assert resp.status_code == 200 and resp.headers["content-type"] == XLSX
    wb = load_workbook(io.BytesIO(resp.content))
    assert wb.sheetnames == ["Операции", "Справочники", "Как заполнять"]
    assert [c.value for c in wb["Операции"][1]] == OP_HEADER
    codes = [r[0].value for r in wb["Справочники"].iter_rows(min_row=2) if r[0].value]
    assert "60.01" in codes and "60" not in codes and "00" not in codes  # только счета без субсчетов
    # пустой шаблон загружается без ошибок
    result = upload(client, env, "operations", resp.content).json()
    assert result["rows"] == 0 and result["errors"] == []


def test_import_permissions(client, env):
    raw = xlsx(OP_HEADER, [op_row(**PAYMENT)])
    ceo = env["headers_for"](env["ceo"])
    other = env["headers_for"](env["other_fd"])
    assert upload(client, env, "operations", raw, h=ceo).status_code == 403
    assert upload(client, env, "counterparties", raw, h=ceo).status_code == 403
    assert upload(client, env, "operations", raw, h=other).status_code == 404
    assert client.get(f"{env['o']}/import/templates/operations.xlsx", headers=ceo).status_code == 403
    assert client.get(f"{env['o']}/import/templates/counterparties.xlsx", headers=other).status_code == 404


# --- Контрагенты ---

CP_HEADER = ["Наименование*", "Полное наименование", "Вид", "ИНН", "КПП", "Покупатель", "Поставщик"]


def test_counterparties_import(client, env):
    rows = [
        ["ООО Ромашка", 'ООО "Ромашка"', "", 7704444444, "770401001", "да", ""],
        ["ИП Иванов", "", "ИП", "770555555555", "", "", "да"],
        ["Петров П.П.", "", "Физлицо", "", "", "нет", "нет"],
        ["ООО Поставщик дубль", "", "", "7702222222", "", "", ""],  # такой ИНН уже есть
        ["Плохой ИНН", "", "", "12345", "", "", ""],
        ["ИП с ИНН организации", "", "ИП", "7706666666", "", "", ""],
        ["", "Без краткого имени", "", "7707777777", "", "", ""],
        ["Флаг", "", "", "7708888888", "", "может быть", ""],
    ]
    resp = upload(client, env, "counterparties", xlsx(CP_HEADER, rows))
    assert resp.status_code == 200, resp.text
    result = resp.json()
    assert result["created"] == ["ООО Ромашка", "ИП Иванов", "Петров П.П."]
    assert result["skipped_existing"] == 1
    assert [e["row"] for e in result["errors"]] == [6, 7, 8, 9]

    by_name = {c["name"]: c for c in client.get(f"{env['o']}/counterparties", headers=env["h"]).json()}
    romashka = by_name["ООО Ромашка"]
    assert (romashka["kind"], romashka["inn"], romashka["kpp"], romashka["is_customer"]) == ("legal", "7704444444", "770401001", True)
    assert (by_name["ИП Иванов"]["kind"], by_name["ИП Иванов"]["is_supplier"]) == ("entrepreneur", True)
    assert by_name["Петров П.П."]["kind"] == "person" and by_name["Петров П.П."]["inn"] is None

    again = upload(client, env, "counterparties", xlsx(CP_HEADER, rows)).json()
    assert again["created"] == [] and again["skipped_existing"] == 4


def test_counterparties_template(client, env):
    resp = client.get(f"{env['o']}/import/templates/counterparties.xlsx", headers=env["h"])
    assert resp.status_code == 200
    wb = load_workbook(io.BytesIO(resp.content))
    assert [c.value for c in wb.active[1]] == CP_HEADER
