"""ОСВ, разрез по аналитике, расшифровка до проводок, выгрузка в Excel (docs/TZ.md, раздел 12)."""

import io

import pytest
from openpyxl import load_workbook

from tests.test_money_page import env, op, pay, receive  # noqa: F401 — общая организация с данными

OCT = {"date_from": "2026-10-01", "date_to": "2026-10-31"}


def cost(env, amount, cfo=None):
    """Аренда магазина: Дт 26 — Кт 60.01."""
    return {"debit_account_id": env["acc"]["26"], "credit_account_id": env["acc"]["60.01"], "amount": amount,
            "debit": {"cfo": cfo or env["shop"]["id"], "pl_item": env["art"]["Аренда"]},
            "credit": {"counterparty": env["cp"]["id"], "contract": env["contract"]["id"]}}


@pytest.fixture
def data(client, env):  # noqa: F811
    op(client, env, [receive(env, "10000")], op_date="2026-09-10")
    op(client, env, [pay(env, "3000")], op_date="2026-10-05")
    op(client, env, [cost(env, "2000")], op_date="2026-10-20")
    op(client, env, [pay(env, "500")], op_date="2026-10-25", post=False)  # черновик
    return env


def get(client, env, path, h=None, **params):
    resp = client.get(f"{env['o']}/reports/{path}", headers=h or env["h"], params={**OCT, **params})
    assert resp.status_code == 200, resp.text
    return resp.json()


def balances(row):
    keys = ["opening_debit", "opening_credit", "debit", "credit", "closing_debit", "closing_credit"]
    return [row[k] for k in keys]


def test_osv(client, data):
    report = get(client, data, "osv")
    rows = {r["code"]: r for r in report["rows"]}
    assert list(rows) == ["26", "51", "60", "60.01"]  # только счета с движением, по порядку кодов
    assert balances(rows["51"]) == ["10000.00", "0.00", "0.00", "3000.00", "7000.00", "0.00"]
    assert balances(rows["60.01"]) == ["0.00", "10000.00", "3000.00", "2000.00", "0.00", "9000.00"]
    assert balances(rows["60"]) == balances(rows["60.01"])
    assert rows["60.01"]["parent_id"] == rows["60"]["account_id"]
    assert balances(rows["26"]) == ["0.00", "0.00", "2000.00", "0.00", "2000.00", "0.00"]
    # итог Дт = итогу Кт; родитель 60 не считается дважды
    assert balances(report["total"]) == ["10000.00", "10000.00", "5000.00", "5000.00", "9000.00", "9000.00"]


def test_osv_with_drafts_and_filters(client, data, db):
    rows = {r["code"]: r for r in get(client, data, "osv", drafts=True)["rows"]}
    assert rows["51"]["credit"] == "3500.00"

    # ЦФО: только проводки с ЦФО из ветки — счёт 26
    rows = {r["code"]: r for r in get(client, data, "osv", cfo_id=data["shop"]["id"])["rows"]}
    assert list(rows) == ["26"]

    # другое юрлицо организации — пусто
    from app.models import LegalEntity

    le2 = LegalEntity(organization_id=data["org"].id, full_name="ИП Б", short_name="Б", inn="770111111111")
    db.add(le2)
    db.commit()
    assert get(client, data, "osv", legal_entity_id=le2.id)["rows"] == []
    assert len(get(client, data, "osv", legal_entity_id=data["le"].id)["rows"]) == 4


def test_cfo_branch_includes_subordinates(client, data):
    child = client.post(f"{data['o']}/cfos", headers=data["h"],
                        json={"name": "Отдел", "parent_id": data["shop"]["id"]}).json()
    op(client, data, [cost(data, "700", cfo=child["id"])], op_date="2026-10-21")
    shop = {r["code"]: r for r in get(client, data, "osv", cfo_id=data["shop"]["id"])["rows"]}
    assert shop["26"]["debit"] == "2700.00"
    only_child = {r["code"]: r for r in get(client, data, "osv", cfo_id=child["id"])["rows"]}
    assert only_child["26"]["debit"] == "700.00"


def test_osv_details_by_analytics(client, data):
    acc = data["acc"]
    details = get(client, data, "osv/details", account_id=acc["60.01"])
    assert details["kind"] == "counterparty"
    assert [k["code"] for k in details["kinds"]] == ["counterparty", "contract"]
    (row,) = details["rows"]
    assert row["name"] == "Поставщик" and row["value"] == data["cp"]["id"]
    assert balances(row) == ["0.00", "10000.00", "3000.00", "2000.00", "0.00", "9000.00"]

    by_contract = get(client, data, "osv/details", account_id=acc["60.01"], kind="contract")
    assert by_contract["rows"][0]["name"] == "№ 1"

    bank = get(client, data, "osv/details", account_id=acc["51"], kind="bank_account")
    assert bank["rows"][0]["name"] == "Основной ··0001"

    resp = client.get(f"{data['o']}/reports/osv/details", headers=data["h"],
                      params={**OCT, "account_id": acc["51"], "kind": "employee"})
    assert resp.status_code == 422


def drill(client, env, **params):
    return get(client, env, "lines", **params)


def test_drill_matches_report(client, data):
    acc = data["acc"]
    osv = {r["code"]: r for r in get(client, data, "osv")["rows"]}

    turnover = drill(client, data, account_id=acc["51"], side="credit")
    assert turnover["count"] == 1 and turnover["credit"] == osv["51"]["credit"]
    line = turnover["rows"][0]
    assert (line["account"], line["corr_account"], line["credit"]) == ("51", "60.01", "3000.00")
    assert "Основной" in line["analytics"] and "Оплата поставщикам" in line["analytics"]

    # сальдо на начало по родительскому счёту 60 — проводки субсчетов
    opening = drill(client, data, account_id=acc["60"], period="opening")
    assert opening["credit"] == osv["60"]["opening_credit"] and opening["debit"] == "0.00"

    # сальдо на конец = разность расшифровки
    closing = drill(client, data, account_id=acc["60.01"], period="closing")
    assert closing["credit"] == "12000.00" and closing["debit"] == "3000.00"  # 9000 Кт

    # ячейка разреза по аналитике
    by_cp = drill(client, data, account_id=acc["60.01"], side="debit", a=[f"counterparty:{data['cp']['id']}"])
    assert by_cp["debit"] == "3000.00"
    assert drill(client, data, account_id=acc["60.01"], a=["counterparty:none"])["count"] == 0

    # черновик — только с флажком
    assert drill(client, data, account_id=acc["51"], side="credit", drafts=True)["credit"] == "3500.00"

    resp = client.get(f"{data['o']}/reports/lines", headers=data["h"], params={**OCT, "a": ["bad:1"]})
    assert resp.status_code == 422


def test_reports_access(client, data):
    ceo = data["headers_for"](data["ceo"])
    other = data["headers_for"](data["other_fd"])
    assert len(get(client, data, "osv", h=ceo)["rows"]) == 4  # CEO видит отчёты
    for path in ("osv", "lines", "osv.xlsx"):
        assert client.get(f"{data['o']}/reports/{path}", headers=other, params=OCT).status_code == 404
    # ЦФО чужой организации
    foreign_cfo = client.post(f"/api/organizations/{data['other'].id}/cfos", headers=other, json={"name": "Чужой"}).json()
    resp = client.get(f"{data['o']}/reports/osv", headers=data["h"], params={**OCT, "cfo_id": foreign_cfo["id"]})
    assert resp.status_code == 404
    resp = client.get(f"{data['o']}/reports/osv", headers=data["h"],
                      params={"date_from": "2026-10-31", "date_to": "2026-10-01"})
    assert resp.status_code == 422


def test_osv_xlsx(client, data):
    resp = client.get(f"{data['o']}/reports/osv.xlsx", headers=data["h"], params=OCT)
    assert resp.status_code == 200
    ws = load_workbook(io.BytesIO(resp.content)).active
    values = [[c.value for c in row] for row in ws.iter_rows()]
    assert values[0][0] == "Оборотно-сальдовая ведомость"
    assert "01.10.2026 — 31.10.2026" in values[1][0]
    by_code = {row[0]: row for row in values[4:]}
    assert by_code["51"][2] == 10000 and by_code["51"][5] == 3000
    assert by_code["Итого"][6] == 9000 and by_code["Итого"][7] == 9000
    indents = {row[0].value: row[0].alignment.indent for row in ws.iter_rows(min_row=4)}
    assert (indents["60"], indents["60.01"]) == (0, 1)  # субсчёт — с отступом под счётом
