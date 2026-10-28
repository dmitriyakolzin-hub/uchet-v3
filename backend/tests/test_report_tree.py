"""Отчёты деревом статей: группа — сумма подстатей, расшифровка группы — вся ветка (docs/TZ.md, раздел 19, «Детали 9.3»)."""

from tests.test_closing import env, operation, rent  # noqa: F401

MARCH = {"date_from": "2026-03-01", "date_to": "2026-03-31"}


def child(client, env, parent_id, name, code):
    resp = client.post(f"{env['o']}/articles", headers=env["h"], json={
        "name": name, "code": code, "parent_id": parent_id, "is_pl": True, "pl_section": "admin"})
    assert resp.status_code == 201, resp.text
    return resp.json()


def post(client, env, article_id, amount, op_date="2026-03-10"):
    body = rent(env, amount)
    body["debit"]["pl_item"] = article_id
    resp = client.post(f"{env['o']}/operations?post=true", headers=env["h"], json={
        "legal_entity_id": env["le"].id, "op_date": op_date, "postings": [body]})
    assert resp.status_code == 201, resp.text


def test_pl_tree_and_drill(client, env):
    group = env["art"]["Аренда"]
    operation(client, env, "2026-03-05")  # 50 000 — ещё на самой «Аренде»
    variable = child(client, env, group, "переменная часть", "2.3.2.")
    gas = child(client, env, variable["id"], "газ", "2.3.2.1.")
    water = child(client, env, variable["id"], "вода", "2.3.2.2.")
    post(client, env, gas["id"], "1000")
    post(client, env, water["id"], "300")

    data = client.get(f"{env['o']}/reports/pl", headers=env["h"], params=MARCH).json()
    rows = {r["key"]: r for r in data["rows"]}
    top = rows[f"artg:{group}"]
    assert (top["kind"], top["values"]["total"], top["parent"], top["level"]) == ("item_group", "-51300.00", None, 1)
    var = rows[f"artg:{variable['id']}"]
    assert (var["title"], var["values"]["total"], var["parent"], var["level"]) == \
        ("2.3.2. переменная часть", "-1300.00", f"artg:{group}", 2)
    assert rows[f"art:{gas['id']}"]["parent"] == f"artg:{variable['id']}"
    own = rows[f"art:{group}"]
    assert (own["title"], own["values"]["total"], own["parent"]) == ("Аренда (без подстатьи)", "-50000.00", f"artg:{group}")
    order = [r["key"] for r in data["rows"]]
    assert order.index(f"art:{water['id']}") > order.index(f"art:{gas['id']}")  # по коду

    # расшифровка группы — проводки всей ветки
    lines = client.get(f"{env['o']}/reports/lines", headers=env["h"], params={
        **MARCH, "period": "turnover", "a_tree": top["drill"]["a_tree"],
        "exclude_corr_account_id": top["drill"]["exclude_corr_ids"]}).json()
    assert (lines["debit"], lines["count"]) == ("51300.00", 3)
    lines = client.get(f"{env['o']}/reports/lines", headers=env["h"], params={
        **MARCH, "period": "turnover", "a_tree": var["drill"]["a_tree"]}).json()
    assert lines["debit"] == "1300.00"
    other = client.get("/api/organizations/2/articles", headers=env["other"]).json()[0]
    assert client.get(f"{env['o']}/reports/lines", headers=env["h"], params={
        **MARCH, "a_tree": f"pl_item:{other['id']}"}).status_code == 404

    # план/факт — те же строки с родителями
    data = client.get(f"{env['o']}/reports/plan-fact", headers=env["h"], params={**MARCH, "kind": "pl"}).json()
    rows = {r["key"]: r for r in data["rows"]}
    assert rows[f"art:{gas['id']}"]["parent"] == f"artg:{variable['id']}"
    assert rows[f"artg:{group}"]["values"]["total"]["fact"] == "-51300.00"

    # Excel: подстатьи сворачиваются
    from io import BytesIO

    from openpyxl import load_workbook

    xlsx = client.get(f"{env['o']}/reports/pl.xlsx", headers=env["h"], params=MARCH)
    ws = load_workbook(BytesIO(xlsx.content)).active
    levels = {ws.cell(row=r, column=1).value: ws.row_dimensions[r].outline_level for r in range(1, ws.max_row + 1)}
    assert levels["2.3.2.1. газ"] == 3
