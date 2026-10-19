"""Управленческий баланс: строки, развёрнутое сальдо расчётов, прибыль периода, расшифровка (docs/TZ.md, раздел 12)."""

import io

from openpyxl import load_workbook

from tests.test_money_page import env, op, opening, pay, set_opening_date  # noqa: F401
from tests.test_reports_pl import cost_side, p, settle

PERIOD = {"date_from": "2026-10-01", "date_to": "2026-10-31"}


def setup_month(client, env, balanced=True):
    assert set_opening_date(client, env, "2026-09-30").status_code == 200
    bank = {"bank_account": env["main"]["id"], "cash_flow_item": env["art"]["Прочие поступления"]}
    postings = [opening(env, "10000", dt="51", ct="00", debit=bank)]
    if balanced:
        postings.append(opening(env, "10000", dt="00", ct="80"))
    op(client, env, postings, op_date="2026-09-30", is_opening=True)
    sale = "Выручка от продаж"
    op(client, env, [p(env, "62.01", "90.01", "12000", None, settle(env), cost_side(env, sale))], op_date="2026-10-05")
    op(client, env, [pay(env, "3000")], op_date="2026-10-06")
    op(client, env, [p(env, "44", "60.01", "1000", None, cost_side(env, "Аренда"), settle(env))], op_date="2026-10-20")
    op(client, env, [p(env, "99", "68.04", "500", None, {"pl_item": env["art"]["Налог на прибыль"]})],
       op_date="2026-10-31")


def balance(client, env, h=None, **params):
    resp = client.get(f"{env['o']}/reports/balance", headers=h or env["h"], params={**PERIOD, **params})
    assert resp.status_code == 200, resp.text
    data = resp.json()
    data["by_key"] = {r["key"]: r for r in data["rows"]}
    return data


def vals(row):
    return [row["values"]["opening"], row["values"]["closing"]]


def test_balance_lines_and_totals(client, env):  # noqa: F811
    setup_month(client, env)
    data = balance(client, env)
    assert [c["title"] for c in data["columns"]] == ["На 30.09.2026", "На 31.10.2026"]
    rows = data["by_key"]
    assert vals(rows["money"]) == ["10000.00", "7000.00"]
    assert vals(rows["customers_dr"]) == ["0.00", "12000.00"]
    assert vals(rows["suppliers_dr"]) == ["0.00", "2000.00"]  # аванс 3000 минус аренда 1000
    assert vals(rows["total_assets"]) == ["10000.00", "21000.00"]
    assert vals(rows["share_capital"]) == ["10000.00", "10000.00"]
    assert vals(rows["current_profit"]) == ["0.00", "10500.00"]  # совпадает с чистой прибылью ОПиУ
    assert vals(rows["taxes_cr"]) == ["0.00", "500.00"]
    assert vals(rows["capital"]) == ["10000.00", "20500.00"]
    assert vals(rows["liabilities"]) == ["0.00", "500.00"]
    assert vals(rows["total_passive"]) == vals(rows["total_assets"])
    assert "unbalanced" not in rows and "suppliers_cr" not in rows  # пустые строки не показываются

    net = client.get(f"{env['o']}/reports/pl", headers=env["h"], params=PERIOD).json()
    assert {r["key"]: r for r in net["rows"]}["net"]["values"]["total"] == "10500.00"


def test_unbalanced_opening_shown(client, env):  # noqa: F811
    setup_month(client, env, balanced=False)
    rows = balance(client, env)["by_key"]
    assert vals(rows["unbalanced"]) == ["10000.00", "10000.00"]
    assert vals(rows["total_passive"]) == vals(rows["total_assets"])


def drill(client, env, **params):
    resp = client.get(f"{env['o']}/reports/lines", headers=env["h"], params=params)
    assert resp.status_code == 200, resp.text
    return resp.json()


def test_balance_drill(client, env):  # noqa: F811
    setup_month(client, env)
    rows = balance(client, env)["by_key"]
    money = rows["money"]["drill"]
    lines = drill(client, env, date_from="2026-10-31", date_to="2026-10-31", period=money["period"],
                  account_id=money["account_id"])
    assert (lines["debit"], lines["credit"]) == ("10000.00", "3000.00")  # 7000

    assert rows["suppliers_dr"]["drill"] == {"details": "suppliers_dr"}
    resp = client.get(f"{env['o']}/reports/balance/details", headers=env["h"],
                      params={"date_from": "2026-10-31", "date_to": "2026-10-31", "line": "suppliers_dr"})
    details = resp.json()
    assert details["total"] == "2000.00" and details["title"] == "Авансы, выданные поставщикам"
    (group,) = details["rows"]
    assert (group["account"], group["name"], group["amount"]) == ("60.01", "Поставщик, № 1", "2000.00")
    a = [f"{k}:{'none' if v is None else v}" for k, v in group["analytics"].items()]
    lines = drill(client, env, date_from="2026-10-31", date_to="2026-10-31", period="closing",
                  account_id=group["account_id"], a=a)
    assert (lines["debit"], lines["credit"]) == ("3000.00", "1000.00")

    bad = client.get(f"{env['o']}/reports/balance/details", headers=env["h"], params={**PERIOD, "line": "nope"})
    assert bad.status_code == 422


def test_balance_xlsx_and_access(client, env):  # noqa: F811
    setup_month(client, env)
    resp = client.get(f"{env['o']}/reports/balance.xlsx", headers=env["h"], params=PERIOD)
    assert resp.status_code == 200
    ws = load_workbook(io.BytesIO(resp.content)).active
    table = {row[0]: row[1:] for row in ws.iter_rows(min_row=4, values_only=True)}
    assert table["Итого актив"] == (10000, 21000) and table["Итого пассив"] == (10000, 21000)

    ceo = env["headers_for"](env["ceo"])
    assert balance(client, env, h=ceo)["by_key"]["total_assets"]["values"]["closing"] == "21000.00"
    other = env["headers_for"](env["other_fd"])
    assert client.get(f"{env['o']}/reports/balance", headers=other, params=PERIOD).status_code == 404
    assert client.get(f"{env['o']}/reports/balance/details", headers=other,
                      params={**PERIOD, "line": "money"}).status_code == 404
