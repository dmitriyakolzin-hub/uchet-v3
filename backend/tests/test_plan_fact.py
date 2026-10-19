"""План/факт: актуальный утверждённый бюджет по месяцам, выбранная версия, расшифровка (docs/TZ.md, раздел 13)."""

import io

import pytest
from openpyxl import load_workbook

from tests.test_plan_grid import env as grid_env  # noqa: F401
from tests.test_plan_grid import put
from tests.test_users import admin  # noqa: F401

PERIOD = {"date_from": "2027-01-01", "date_to": "2027-02-28"}


@pytest.fixture
def env(client, grid_env, db):  # noqa: F811
    env = grid_env
    h, o = env["h"], env["o"]
    rub = client.get("/api/currencies", headers=h).json()[0]["id"]
    cp = client.post(f"{o}/counterparties", headers=h, json={"name": "Арендодатель", "kind": "legal",
                                                            "inn": "7702222222"}).json()
    contract = client.post(f"{o}/contracts", headers=h, json={"counterparty_id": cp["id"], "legal_entity_id": env["le"].id,
                                                             "currency_id": rub, "kind": "supplier", "number": "А-1"}).json()
    env.update(cp=cp, contract=contract)
    # годовой бюджет: аренда магазина 1000 в январе и феврале, оплата поставщикам 500 в январе
    rent, pay = env["art"]["Аренда"], env["art"]["Оплата поставщикам"]
    put(client, env, [{"article_id": rent, "month": 1, "amount": "1000"}, {"article_id": rent, "month": 2, "amount": "1000"}],
        cfo="shop")
    put(client, env, [{"article_id": pay, "month": 1, "amount": "500"}], kind="cf")
    assert client.post(f"/api/budget-versions/{env['version']['id']}/approve", headers=h).status_code == 200
    # уточнённый бюджет на январь: аренда 1200
    jan = client.post(f"{o}/budgets", headers=h, json={"kind": "month", "year": 2027, "period": 1}).json()
    env["jan_version"] = jan["versions"][0]
    put(client, env, [{"article_id": rent, "month": 1, "amount": "1200"}], cfo="shop", version=env["jan_version"])
    assert client.post(f"/api/budget-versions/{env['jan_version']['id']}/approve", headers=h).status_code == 200
    # факт: аренда 900 в январе и 1100 в феврале
    acc = {v: k for k, v in env["acc"].items()}
    for date, amount in (("2027-01-15", "900"), ("2027-02-10", "1100")):
        resp = client.post(f"{o}/operations?post=true", headers=h, json={
            "legal_entity_id": env["le"].id, "op_date": date, "postings": [{
                "debit_account_id": acc["44"], "credit_account_id": acc["60.01"], "amount": amount,
                "debit": {"cfo": env["shop"]["id"], "pl_item": rent},
                "credit": {"counterparty": cp["id"], "contract": contract["id"]}}]})
        assert resp.status_code == 201, resp.text
    return env


def report(client, env, h=None, **params):
    resp = client.get(f"{env['o']}/reports/plan-fact", headers=h or env["h"], params={**PERIOD, **params})
    assert resp.status_code == 200, resp.text
    data = resp.json()
    data["by_key"] = {r["key"]: r for r in data["rows"]}
    return data


def test_actual_budget_by_month(client, env):
    data = report(client, env)
    cols = {c["key"]: c for c in data["columns"]}
    assert cols["2027-01"]["version_id"] == env["jan_version"]["id"]  # месячный бюджет важнее годового
    assert cols["2027-02"]["version_id"] == env["version"]["id"]
    assert cols["2027-01"]["version_title"] == "Бюджет январь 2027: Версия 1"
    rent = data["by_key"][f"art:{env['art']['Аренда']}"]["values"]
    assert rent["2027-01"] == {"plan": "-1200.00", "fact": "-900.00", "dev": "300.00", "pct": "75.0"}
    assert rent["2027-02"] == {"plan": "-1000.00", "fact": "-1100.00", "dev": "-100.00", "pct": "110.0"}
    assert rent["total"] == {"plan": "-2200.00", "fact": "-2000.00", "dev": "200.00", "pct": "90.9"}
    assert data["by_key"]["net"]["values"]["total"]["plan"] == "-2200.00"
    # статьи без плана и факта не показываются, разделы ОПиУ — всегда
    assert f"art:{env['art']['Связь и IT']}" not in data["by_key"] and "revenue" in data["by_key"]


def test_selected_version(client, env):
    data = report(client, env, version_id=env["version"]["id"])
    rent = data["by_key"][f"art:{env['art']['Аренда']}"]["values"]
    assert rent["2027-01"]["plan"] == "-1000.00"
    month_only = report(client, env, version_id=env["jan_version"]["id"])
    assert month_only["by_key"][f"art:{env['art']['Аренда']}"]["values"]["2027-02"]["plan"] == "0.00"
    assert {c["key"]: c["version_id"] for c in month_only["columns"]}["2027-02"] is None


def test_plan_drill_matches_cell(client, env):
    row = report(client, env)["by_key"][f"art:{env['art']['Аренда']}"]
    params = {"date_from": "2027-01-01", "date_to": "2027-01-31", "period": "turnover", "a": row["drill"]["a"]}
    plan = client.get(f"{env['o']}/reports/lines", headers=env["h"],
                      params={**params, "plan_version_id": env["jan_version"]["id"]}).json()
    assert (plan["debit"], plan["count"]) == ("1200.00", 1)
    fact = client.get(f"{env['o']}/reports/lines", headers=env["h"], params=params).json()
    assert fact["debit"] == "900.00"


def test_cash_plan_fact(client, env):
    data = report(client, env, kind="cf")
    keys = [r["key"] for r in data["rows"]]
    assert "opening" not in keys and "closing" not in keys
    pay = data["by_key"][f"art:{env['art']['Оплата поставщикам']}"]["values"]
    assert pay["2027-01"] == {"plan": "-500.00", "fact": "0.00", "dev": "500.00", "pct": "0.0"}
    assert "title:investing" not in keys  # пустые виды деятельности убраны


def test_plan_fact_access_and_excel(client, env, make_org, headers_for, make_user):
    resp = client.get(f"{env['o']}/reports/plan-fact.xlsx", headers=env["h"], params=PERIOD)
    ws = load_workbook(io.BytesIO(resp.content)).active
    table = {row[0]: row[1:] for row in ws.iter_rows(min_row=4, values_only=True)}
    assert table["Аренда"][:3] == (-1200, -900, 300)
    assert table["Аренда"][-1] == 90.9

    assert report(client, env, h=env["ceo"])["by_key"]["net"]["values"]["total"]["fact"] == "-2000.00"
    # руководитель ЦФО — план и факт только своей ветки (магазин — в его ветке «Торговля»)
    head = report(client, env, h=env["head"])
    assert head["by_key"]["net"]["values"]["total"]["plan"] == "-2200.00"
    other = make_user(login="stranger")
    foreign = make_org("Клиент Б", members=[(other, "finance_director")])
    resp = client.get(f"/api/organizations/{foreign.id}/reports/plan-fact", headers=headers_for(other),
                      params={**PERIOD, "version_id": env["version"]["id"]})
    assert resp.status_code == 404
