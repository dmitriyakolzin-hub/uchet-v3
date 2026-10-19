"""Отчёт ДДС: месяцы, виды деятельности, переводы, остатки, расшифровка, Excel (docs/TZ.md, раздел 12)."""

import io

from openpyxl import load_workbook

from tests.test_money_page import env, money, op, opening, pay, receive, set_opening_date, transfer  # noqa: F401

PERIOD = {"date_from": "2026-10-01", "date_to": "2026-11-30"}


def dds(client, env, h=None, **params):
    resp = client.get(f"{env['o']}/reports/dds", headers=h or env["h"], params={**PERIOD, **params})
    assert resp.status_code == 200, resp.text
    data = resp.json()
    data["by_key"] = {r["key"]: r for r in data["rows"]}
    return data


def values(row):
    return [row["values"]["2026-10"], row["values"]["2026-11"], row["values"]["total"]]


def month_data(client, env):
    op(client, env, [receive(env, "10000")], op_date="2026-09-10")
    op(client, env, [pay(env, "3000")], op_date="2026-10-05")
    op(client, env, [transfer(env, "1000")], op_date="2026-10-07")
    op(client, env, [pay(env, "500")], op_date="2026-11-03")
    op(client, env, [pay(env, "200")], op_date="2026-10-25", post=False)


def test_dds_by_months(client, env):  # noqa: F811
    month_data(client, env)
    data = dds(client, env)
    assert [c["key"] for c in data["columns"]] == ["2026-10", "2026-11", "total"]
    assert data["columns"][0]["title"] == "Октябрь 2026"
    rows = data["by_key"]
    supplier = rows[f"art:{env['art']['Оплата поставщикам']}"]
    assert values(supplier) == ["-3000.00", "-500.00", "-3500.00"]
    assert supplier["level"] == 2
    assert values(rows["grp:operating:outflow"]) == ["-3000.00", "-500.00", "-3500.00"]
    assert values(rows["act:operating"]) == ["-3000.00", "-500.00", "-3500.00"]
    assert values(rows["net"]) == ["-3000.00", "-500.00", "-3500.00"]
    assert values(rows["opening"]) == ["10000.00", "7000.00", "10000.00"]
    assert values(rows["closing"]) == ["7000.00", "6500.00", "6500.00"]
    # перевод между своими счетами в отчёт не попадает
    assert f"art:{env['art']['Перевод между своими счетами']}" not in rows and "moves" not in rows
    # остаток совпадает со страницей «Деньги»
    assert money(client, env, "2026-11-01", "2026-11-30")["closing"] == "6500.00"

    with_drafts = dds(client, env, drafts=True)["by_key"]
    assert values(with_drafts[f"art:{env['art']['Оплата поставщикам']}"])[0] == "-3200.00"


def drill(client, env, **params):
    resp = client.get(f"{env['o']}/reports/lines", headers=env["h"], params={**PERIOD, **params})
    assert resp.status_code == 200, resp.text
    return resp.json()


def test_dds_drill_matches_cells(client, env):  # noqa: F811
    month_data(client, env)
    data = dds(client, env)
    money_ids = data["money_account_ids"]
    october = data["columns"][0]
    supplier = data["by_key"][f"art:{env['art']['Оплата поставщикам']}"]
    d = supplier["drill"]
    lines = drill(client, env, date_from=october["date_from"], date_to=october["date_to"], period=d["period"],
                  account_id=money_ids, exclude_corr_account_id=money_ids, opening=False, a=d["a"])
    assert lines["count"] == 1 and lines["credit"] == "3000.00" and lines["debit"] == "0.00"

    closing = drill(client, env, date_from="2026-11-01", date_to="2026-11-30", period="closing", account_id=money_ids)
    assert closing["debit"] == "11000.00" and closing["credit"] == "4500.00"  # 6500 с переводом в обе стороны


def test_cash_moves_under_cfo_filter(client, env):  # noqa: F811
    cash_in = {"debit_account_id": env["acc"]["50"], "credit_account_id": env["acc"]["51"], "amount": "100",
               "debit": {"cfo": env["shop"]["id"], "cash_flow_item": env["art"]["Перевод между своими счетами"]},
               "credit": {"bank_account": env["main"]["id"],
                          "cash_flow_item": env["art"]["Перевод между своими счетами"]}}
    op(client, env, [cash_in], op_date="2026-10-10")
    # по всей организации — перемещение внутри денег, потока нет
    rows = dds(client, env)["by_key"]
    assert "moves" not in rows and values(rows["closing"])[2] == "0.00"
    # по ЦФО видна только касса: деньги «пришли» перемещением, остатки сходятся
    rows = dds(client, env, cfo_id=env["shop"]["id"])["by_key"]
    assert values(rows["moves"]) == ["100.00", "0.00", "100.00"]
    assert values(rows["closing"]) == ["100.00", "100.00", "100.00"]


def test_opening_entries_row(client, env):  # noqa: F811
    assert set_opening_date(client, env, "2026-09-30").status_code == 200
    money_in = opening(env, "5000", dt="51", ct="00",
                       debit={"bank_account": env["main"]["id"], "cash_flow_item": env["art"]["Прочие поступления"]})
    op(client, env, [money_in], op_date="2026-09-30", is_opening=True)
    op(client, env, [pay(env, "1000")], op_date="2026-10-02")

    rows = dds(client, env, date_from="2026-09-01", date_to="2026-10-31")["by_key"]
    assert rows["opening_entries"]["values"] == {"2026-09": "5000.00", "2026-10": "0.00", "total": "5000.00"}
    # ввод остатков — не поступление по статье
    assert f"art:{env['art']['Прочие поступления']}" not in rows
    assert rows["closing"]["values"]["total"] == "4000.00"
    lines = drill(client, env, date_from="2026-09-01", date_to="2026-09-30", period="turnover",
                  account_id=dds(client, env)["money_account_ids"], opening=True)
    assert lines["debit"] == "5000.00"

    # отчёт с октября: остатки уже в сальдо на начало
    rows = dds(client, env, date_from="2026-10-01", date_to="2026-10-31")["by_key"]
    assert rows["opening"]["values"]["total"] == "5000.00" and "opening_entries" not in rows


def test_dds_xlsx_and_access(client, env):  # noqa: F811
    month_data(client, env)
    resp = client.get(f"{env['o']}/reports/dds.xlsx", headers=env["h"], params=PERIOD)
    assert resp.status_code == 200
    ws = load_workbook(io.BytesIO(resp.content)).active
    table = {row[0]: row[1:] for row in ws.iter_rows(min_row=4, values_only=True)}
    assert table["Статья"] == ("Октябрь 2026", "Ноябрь 2026", "Итого")
    assert table["Остаток денег на конец"] == (7000, 6500, 6500)
    assert table["Оплата поставщикам"] == (-3000, -500, -3500)

    ceo = env["headers_for"](env["ceo"])
    assert dds(client, env, h=ceo)["by_key"]["closing"]["values"]["total"] == "6500.00"
    other = env["headers_for"](env["other_fd"])
    assert client.get(f"{env['o']}/reports/dds", headers=other, params=PERIOD).status_code == 404
