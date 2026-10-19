"""Сотрудники: внутренний №, выгрузка и загрузка из Excel (docs/TZ.md, раздел 17, «Детали 7.1–7.2»)."""

import io

import pytest
from openpyxl import load_workbook

from tests.test_excel_import import XLSX, xlsx

HEADER = ["ФИО*", "Табельный №", "Внутренний №", "ИНН", "Должность", "Юрлицо (ИНН или название)", "ЦФО",
          "Статья затрат", "Дата приёма", "Дата увольнения", "Оклад", "Оклад с даты"]


@pytest.fixture
def env(client, headers_for, make_user, make_org, db):
    from app.models import LegalEntity

    fd, ceo = make_user(login="fd"), make_user(login="ceo")
    org = make_org("Группа А", members=[(fd, "finance_director"), (ceo, "ceo")])
    le = LegalEntity(organization_id=org.id, full_name="ООО А", short_name="А", inn="7701234567")
    db.add(le)
    db.commit()
    h = headers_for(fd)
    o = f"/api/organizations/{org.id}"
    company = client.post(f"{o}/cfos", headers=h, json={"name": "Компания"}).json()
    shop = client.post(f"{o}/cfos", headers=h, json={"name": "Магазин 1", "parent_id": company["id"]}).json()
    return {"o": o, "h": h, "le": le, "shop": shop, "ceo": headers_for(ceo)}


def upload(client, env, rows, h=None):
    return client.post(f"{env['o']}/import/employees", headers={**(h or env["h"]), "Content-Type": XLSX},
                       content=xlsx(HEADER, rows))


def employees(client, env):
    return {e["full_name"]: e for e in client.get(f"{env['o']}/employees", headers=env["h"]).json()}


def test_internal_number_unique(client, env):
    post = lambda body: client.post(f"{env['o']}/employees", headers=env["h"], json=body)  # noqa: E731
    assert post({"full_name": "А", "personnel_number": "1", "internal_number": "A-1"}).status_code == 201
    assert post({"full_name": "Б", "internal_number": "A-1"}).status_code == 409
    assert post({"full_name": "В", "internal_number": ""}).status_code == 201  # пустой не мешает
    assert post({"full_name": "Г"}).json()["internal_number"] == ""


def test_import_creates_and_updates(client, env):
    resp = upload(client, env, [
        ["Иванов Иван", "001", "M-17", "770000000001", "Продавец", "7701234567", "Компания / Магазин 1",
         "Оплата труда — коммерческий персонал", "01.02.2026", None, 60000, None],
        ["Петров Пётр", None, "M-18", None, None, "А", "Магазин 1", None, "2026-03-15", None, "45 000,50", "01.04.2026"],
        ["Без номера", None, None, "123", None, None, None, None, None, None, None, None],  # ИНН не 12 цифр
        ["Чужой ЦФО", "9", None, None, None, None, "Нет такого", None, None, None, None, None],
    ])
    assert resp.status_code == 200, resp.text
    result = resp.json()
    assert result["created"] == ["Иванов Иван", "Петров Пётр"] and result["updated"] == []
    assert [(e["row"], "ИНН" in e["error"] or "ЦФО" in e["error"]) for e in result["errors"]] == [(4, True), (5, True)]

    staff = employees(client, env)
    ivanov = staff["Иванов Иван"]
    assert (ivanov["personnel_number"], ivanov["internal_number"], ivanov["legal_entity_id"], ivanov["cfo_id"],
            ivanov["hired_on"]) == ("001", "M-17", env["le"].id, env["shop"]["id"], "2026-02-01")
    assert ivanov["pl_item_id"] is not None
    salaries = client.get(f"/api/employees/{ivanov['id']}/salaries", headers=env["h"]).json()
    assert [(s["valid_from"], s["amount"]) for s in salaries] == [("2026-02-01", "60000.00")]
    petrov = staff["Петров Пётр"]
    salaries = client.get(f"/api/employees/{petrov['id']}/salaries", headers=env["h"]).json()
    assert [(s["valid_from"], s["amount"]) for s in salaries] == [("2026-04-01", "45000.50")]

    # обновление по внутреннему №: пустые ячейки не стирают, новый оклад — в историю
    resp = upload(client, env, [["Петров Пётр Петрович", None, "M-18", None, "Старший продавец", None, None, None,
                                 None, "31.12.2026", 50000, "01.07.2026"]])
    assert resp.json()["updated"] == ["Петров Пётр Петрович"], resp.text
    petrov = employees(client, env)["Петров Пётр Петрович"]
    assert (petrov["position"], petrov["hired_on"], petrov["fired_on"], petrov["cfo_id"]) == (
        "Старший продавец", "2026-03-15", "2026-12-31", env["shop"]["id"])
    salaries = client.get(f"/api/employees/{petrov['id']}/salaries", headers=env["h"]).json()
    assert len(salaries) == 2

    # конфликт номеров
    resp = upload(client, env, [["Кто-то", "001", "M-18", None, None, None, None, None, None, None, None, None]])
    assert "у разных сотрудников" in resp.json()["errors"][0]["error"]


def test_employees_download_and_access(client, env):
    upload(client, env, [["Иванов Иван", "001", "M-17", None, None, "А", None, None, "01.02.2026", None, 60000, None]])
    resp = client.get(f"{env['o']}/import/employees.xlsx", headers=env["h"])
    assert resp.status_code == 200
    ws = load_workbook(io.BytesIO(resp.content)).active
    assert [c.value for c in ws[1]][:3] == ["ФИО*", "Табельный №", "Внутренний №"]
    row = [c.value for c in ws[2]]
    assert (row[0], row[1], row[2], row[5], row[10]) == ("Иванов Иван", "001", "M-17", "7701234567", 60000)
    # повторная загрузка скачанного файла ничего не ломает
    resp = client.post(f"{env['o']}/import/employees", headers={**env["h"], "Content-Type": XLSX}, content=resp.content)
    assert resp.json()["updated"] == ["Иванов Иван"] and resp.json()["errors"] == []
    assert client.get(f"{env['o']}/import/templates/employees.xlsx", headers=env["h"]).status_code == 200
    assert upload(client, env, [], h=env["ceo"]).status_code == 403
    assert client.get(f"{env['o']}/import/employees.xlsx", headers=env["ceo"]).status_code == 403
