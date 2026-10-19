"""Дебиторская и кредиторская задолженность: колонки по знаку, авансы, раскрытие по договорам (docs/TZ.md, раздел 12)."""

import io

from openpyxl import load_workbook

from tests.test_money_page import env, op, pay  # noqa: F401
from tests.test_reports_pl import cost_side, p, settle

PERIOD = {"date_from": "2026-10-01", "date_to": "2026-10-31"}


def make_data(client, env):
    o, h = env["o"], env["h"]
    buyer = client.post(f"{o}/counterparties", headers=h, json={"name": "Покупатель", "kind": "legal",
                                                                 "inn": "7703333333"}).json()
    rub = client.get("/api/currencies", headers=h).json()[0]["id"]
    contract = client.post(f"{o}/contracts", headers=h, json={
        "counterparty_id": buyer["id"], "legal_entity_id": env["le"].id, "currency_id": rub,
        "kind": "customer", "number": "П-1"}).json()
    sale = {"counterparty": buyer["id"], "contract": contract["id"]}
    op(client, env, [p(env, "62.01", "90.01", "5000", None, sale, cost_side(env, "Выручка от продаж"))],
       op_date="2026-09-15")
    op(client, env, [p(env, "62.01", "90.01", "7000", None, sale, cost_side(env, "Выручка от продаж"))],
       op_date="2026-10-05")
    op(client, env, [pay(env, "3000")], op_date="2026-10-06")  # аванс поставщику
    op(client, env, [p(env, "44", "60.01", "4500", None, cost_side(env, "Аренда"), settle(env))],
       op_date="2026-10-20")  # аренда больше аванса — долг поставщику 1500
    return buyer, contract


def debts(client, env, h=None, **params):
    resp = client.get(f"{env['o']}/reports/debts", headers=h or env["h"], params={**PERIOD, **params})
    assert resp.status_code == 200, resp.text
    return resp.json()


def cols(row):
    return [row[k] for k in ("receivable", "advance_paid", "payable", "advance_received", "opening", "saldo")]


def test_debts_by_counterparty(client, env):  # noqa: F811
    buyer, contract = make_data(client, env)
    data = debts(client, env)
    rows = {r["key"]: r for r in data["rows"]}
    buyer_row = rows[f"cp:{buyer['id']}"]
    assert buyer_row["parent"] is None and buyer_row["title"] == "Покупатель"
    assert cols(buyer_row) == ["12000.00", "0.00", "0.00", "0.00", "5000.00", "12000.00"]
    (child,) = [r for r in data["rows"] if r["parent"] == buyer_row["key"]]
    assert (child["account"], child["title"]) == ("62.01", "№ П-1")
    supplier = rows[f"cp:{env['cp']['id']}"]
    assert cols(supplier) == ["0.00", "0.00", "1500.00", "0.00", "0.00", "-1500.00"]
    assert cols(data["total"]) == ["12000.00", "0.00", "1500.00", "0.00", "5000.00", "10500.00"]
    # покупатели по алфавиту
    parents = [r["title"] for r in data["rows"] if r["parent"] is None]
    assert parents == sorted(parents, key=str.lower)

    # на 10.10 аванс поставщику ещё не закрыт
    early = {r["key"]: r for r in debts(client, env, date_to="2026-10-10")["rows"]}
    assert cols(early[f"cp:{env['cp']['id']}"])[1] == "3000.00"


def test_debts_drill(client, env):  # noqa: F811
    buyer, _ = make_data(client, env)
    rows = {r["key"]: r for r in debts(client, env)["rows"]}
    parent = rows[f"cp:{buyer['id']}"]
    resp = client.get(f"{env['o']}/reports/lines", headers=env["h"], params={
        **PERIOD, "period": "closing", "account_id": parent["drill"]["account_id"], "a": parent["drill"]["a"]})
    lines = resp.json()
    assert (lines["debit"], lines["credit"]) == ("12000.00", "0.00")
    opening = client.get(f"{env['o']}/reports/lines", headers=env["h"], params={
        **PERIOD, "period": "opening", "account_id": parent["drill"]["account_id"], "a": parent["drill"]["a"]}).json()
    assert opening["debit"] == "5000.00"
    child = next(r for r in rows.values() if r["parent"] == f"cp:{env['cp']['id']}")
    lines = client.get(f"{env['o']}/reports/lines", headers=env["h"], params={
        **PERIOD, "period": "closing", "account_id": child["drill"]["account_id"], "a": child["drill"]["a"]}).json()
    assert (lines["debit"], lines["credit"]) == ("3000.00", "4500.00")


def test_debts_xlsx_and_access(client, env):  # noqa: F811
    make_data(client, env)
    resp = client.get(f"{env['o']}/reports/debts.xlsx", headers=env["h"], params=PERIOD)
    ws = load_workbook(io.BytesIO(resp.content)).active
    table = {row[0]: row[1:] for row in ws.iter_rows(min_row=4, values_only=True)}
    assert table["Итого"] == (None, 12000, 0, 1500, 0, 5000, 10500)
    ceo = env["headers_for"](env["ceo"])
    assert len(debts(client, env, h=ceo)["rows"]) == 4
    other = env["headers_for"](env["other_fd"])
    assert client.get(f"{env['o']}/reports/debts", headers=other, params=PERIOD).status_code == 404
